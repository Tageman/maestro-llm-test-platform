from dataclasses import replace
from datetime import UTC, datetime
import json
import threading

import pytest

from maestro_app.application.pytest_generator import GeneratedPytestPackageBuilder
from maestro_app.domain.enums import FixturePhase, FixtureScope, RunMode, WorkbookStatus
from maestro_app.domain.models import DevicePreparation, Run, TestCaseDefinition as Case, WorkbookSnapshot
from maestro_app.infrastructure.pytest_process import ManagedPytestRunner
from maestro_app.infrastructure.run_repository import RunRepository
from unit.test_pytest_lifecycle import _config, _fake_maestro, _fixture, _settings, _step


def shared_package(tmp_path, *, failure=None, startup=None, inconsistent_startup=False):
    maestro = _fake_maestro(tmp_path / "fake-maestro")
    config = _config(maestro)
    fixtures = []
    for index, (key, scope, phase) in enumerate([
        ("S", "session", "setup"), ("M", "module", "setup"), ("F", "function", "setup"),
        ("PS", "session", "teardown"), ("PM", "module", "teardown"), ("PF", "function", "teardown"),
    ], 1):
        fixtures.append(_fixture(key, FixturePhase(phase), index).model_copy(update={"scope": FixtureScope(scope)}))
    cases = [Case(case_id=key, module=module, title=key, pre_fixture_ids=["S", "M", "F"], post_fixture_ids=["PS", "PM", "PF"], priority="P0", steps=_step("点击【按钮】"), expected_results=_step("页面显示【完成】"), source_row=index) for index, (key, module) in enumerate([("A1", "A"), ("A2", "A"), ("B1", "B")], 6)]
    if inconsistent_startup:
        cases[1] = cases[1].model_copy(update={"pre_fixture_ids": ["M", "F"]})
    snapshot = WorkbookSnapshot(workbook_id="wb-shared", source_filename="source.xlsx", source_sha256="a" * 64, imported_at=datetime.now(UTC), status=WorkbookStatus.VALID, fixtures=fixtures, cases=cases, config=config)
    run = Run(run_id="run-shared", generation_id="gen-shared", workbook_id=snapshot.workbook_id, mode=RunMode.REVIEW, selected_case_ids=[case.case_id for case in cases], config_snapshot=config)
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"synthetic source")
    paths = {}
    for key in [fixture.fixture_id for fixture in fixtures] + run.selected_case_ids:
        path = tmp_path / f"{key}.yaml"
        path.write_text("FAIL" if key == failure else "PASS")
        if key == startup:
            path.write_text("appId: example.app\n---\n- launchApp:\n    clearState: false\n")
        paths[key] = path
    repository = RunRepository(tmp_path)
    builder = GeneratedPytestPackageBuilder(repository)
    package = builder.generate(run=run, snapshot=snapshot, flow_paths=paths, source_workbook=source)
    device = DevicePreparation(device_id="emulator-5554", avd_name="Pixel_5", app_id=config.app_id, app_installed=True, boot_completed=True)
    return run, config, repository, builder, package, device


def execute(data, runner, case_id, ordinal=1):
    run, config, repository, builder, package, device = data
    root = repository.run_dir(run.run_id)
    mapping = json.loads((root / "generated/mappings/pytest.json").read_text())
    attempt = root / "attempts" / case_id / f"a{ordinal}"
    package = replace(package, node_ids=(mapping["case_node_ids"][case_id],), junit_path=attempt / "reports/pytest-junit.xml", trace_path=attempt / "reports/flow-trace.jsonl", stdout_path=attempt / "artifacts/pytest.log")
    builder.configure(package, config=config, device=device, run_dir=root, attempt_root=attempt)
    return runner.run(run.run_id, package, timeout_seconds=20, cancel_event=threading.Event()), package


def traces(data):
    root = data[2].run_dir(data[0].run_id)
    return [json.loads(line) for path in root.rglob("flow-trace.jsonl") for line in path.read_text().splitlines()]


def test_true_session_module_function_scopes_share_one_pytest_process(tmp_path):
    data = shared_package(tmp_path)
    runner = ManagedPytestRunner(_settings(tmp_path))
    assert runner.collect(data[4]).succeeded
    try:
        pids = []
        for case_id in ["A1", "A2", "B1"]:
            result, _ = execute(data, runner, case_id)
            assert result.succeeded, result.stdout_tail
            pids.append(runner._shared[data[0].run_id].process.pid)
        assert len(set(pids)) == 1
        runner.close_run(data[0].run_id)
    finally:
        runner.cancel(data[0].run_id)
        runner.close_run(data[0].run_id)
    finished = [event["entity_id"] for event in traces(data) if event["event"] == "finished" and event["entity_type"] in {"case", "fixture"}]
    assert {key: finished.count(key) for key in ["S", "M", "F", "PS", "PM", "PF", "SYS-DEFAULT-LAUNCH"]} == {"S": 1, "M": 2, "F": 3, "PS": 1, "PM": 2, "PF": 3, "SYS-DEFAULT-LAUNCH": 1}
    assert finished.index("A2") < finished.index("PM") < finished.index("B1")
    session_end = next(event for event in traces(data) if event["entity_id"] == "PS" and event["event"] == "finished")
    assert session_end["affected_case_ids"] == ["A1", "A2", "B1"]


def test_retry_rebuilds_full_scopes_then_shares_new_epoch(tmp_path):
    data = shared_package(tmp_path, failure="A1")
    runner = ManagedPytestRunner(_settings(tmp_path))
    try:
        failed, _ = execute(data, runner, "A1")
        assert not failed.succeeded
        (data[2].run_dir(data[0].run_id) / "generated/cases/A1.yaml").write_text("PASS")
        for case_id, ordinal in [("A1", 2), ("A2", 1), ("B1", 1)]:
            result, _ = execute(data, runner, case_id, ordinal)
            assert result.succeeded, result.stdout_tail
    finally:
        runner.close_run(data[0].run_id)
    finished = [event["entity_id"] for event in traces(data) if event["event"] == "finished"]
    assert finished.count("S") == 2
    assert finished.count("M") == 3
    assert finished.count("PS") == 2
    assert finished.count("SYS-DEFAULT-LAUNCH") == 2


def test_failed_shared_setup_is_cached_and_never_runs_dependent_cases(tmp_path):
    data = shared_package(tmp_path, failure="S")
    runner = ManagedPytestRunner(_settings(tmp_path))
    try:
        for case_id in ["A1", "A2", "B1"]:
            result, _ = execute(data, runner, case_id)
            assert not result.succeeded
    finally:
        runner.close_run(data[0].run_id)
    events = traces(data)
    assert sum(event["entity_id"] == "S" and event["event"] == "started" for event in events) == 1
    assert not any(event["entity_type"] == "case" for event in events)
    assert any(event["event"] == "dependency_failed" and event["case_id"] == "A2" for event in events)


def test_failfast_close_flushes_shared_cleanup_without_running_next_case(tmp_path):
    data = shared_package(tmp_path, failure="A1")
    runner = ManagedPytestRunner(_settings(tmp_path))
    result, _ = execute(data, runner, "A1")
    assert not result.succeeded
    closed = runner.close_run(data[0].run_id)
    assert closed[1].succeeded
    events = traces(data)
    assert {event["entity_id"] for event in events if event["phase"] == "teardown" and event["event"] == "finished"} == {"PS", "PM", "PF"}
    assert not any(event["case_id"] in {"A2", "B1"} for event in events)


def test_cancelled_worker_cannot_report_successful_shared_cleanup(tmp_path):
    data = shared_package(tmp_path)
    runner = ManagedPytestRunner(_settings(tmp_path))
    assert execute(data, runner, "A1")[0].succeeded
    runner.cancel(data[0].run_id)
    closed = runner.close_run(data[0].run_id)
    assert not closed[1].succeeded
    assert "无法确认" in closed[1].stdout_tail
    assert runner.close_run(data[0].run_id) is None


@pytest.mark.parametrize("startup", ["F", "M", "A1"])
def test_shorter_scope_startup_is_rejected_before_execution(tmp_path, startup):
    with pytest.raises(ValueError, match="最宽作用域"):
        shared_package(tmp_path, startup=startup)


def test_shared_startup_must_be_consistent_across_cases(tmp_path):
    with pytest.raises(ValueError, match="一致"):
        shared_package(tmp_path, startup="S", inconsistent_startup=True)


def test_session_startup_disables_default_launch_for_entire_epoch(tmp_path):
    data = shared_package(tmp_path, startup="S")
    config = json.loads(data[4].config_path.read_text())
    assert all(not plan["enabled"] for plan in config["startup_plans"].values())
    assert "SYS-DEFAULT-LAUNCH" not in config["flow_paths"]


def test_shared_worker_timeout_retains_unknown_cleanup_evidence(tmp_path):
    data = shared_package(tmp_path)
    runner = ManagedPytestRunner(_settings(tmp_path))
    data[3].configure(data[4], config=data[1], device=data[5], run_dir=data[2].run_dir(data[0].run_id))
    package = replace(data[4], node_ids=(data[4].node_ids[0],))
    result = runner.run(data[0].run_id, package, timeout_seconds=0, cancel_event=threading.Event())
    assert result.timed_out and not result.succeeded
    assert not runner.close_run(data[0].run_id)[1].succeeded


def test_shared_cleanup_failure_lists_all_scope_consumers(tmp_path):
    data = shared_package(tmp_path, failure="PS")
    runner = ManagedPytestRunner(_settings(tmp_path))
    try:
        assert execute(data, runner, "A1")[0].succeeded
        assert execute(data, runner, "A2")[0].succeeded
        assert not execute(data, runner, "B1")[0].succeeded
    finally:
        assert runner.close_run(data[0].run_id)[1].succeeded
    event = next(event for event in traces(data) if event["entity_id"] == "PS" and event["event"] == "finished")
    assert event["succeeded"] is False
    assert event["affected_case_ids"] == ["A1", "A2", "B1"]


@pytest.mark.parametrize("cancel", [False, True])
def test_run_service_shared_cleanup_backpropagation_and_writeback(tmp_path, cancel):
    from openpyxl import load_workbook
    from maestro_app.application.flow_translator import RuleBasedFlowTranslator
    from maestro_app.application.generation_service import GenerationService
    from maestro_app.application.run_service import RunService
    from maestro_app.application.workbook_service import WorkbookService
    from maestro_app.domain.enums import CaseRunState, RunState
    from maestro_app.domain.models import CreateGenerationRequest, CreateRunRequest
    from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
    from maestro_app.infrastructure.generation_repository import GenerationRepository
    from maestro_app.infrastructure.result_workbook import ResultWorkbookWriter
    from maestro_app.infrastructure.workbook_repository import WorkbookRepository
    from integration.test_run_service import PassingSyntaxChecker, PreparedDevice
    from m1_helpers import build_sample_workbook

    maestro = _fake_maestro(tmp_path / "fake-maestro")
    source = build_sample_workbook(tmp_path / "source.xlsx")
    workbook = load_workbook(source)
    for row in workbook["Config"].iter_rows(min_row=6):
        if row[0].value == "maestro_command":
            row[1].value = str(maestro)
    fixtures = workbook["Fixtures"]
    fixtures.cell(6, 4, "会话")
    fixtures.cell(7, 4, "模块")
    fixtures.append(["PS", "清理", "后置", "会话", "", "", 30, "是", "1. 点击【完成】" if cancel else "1. 点击【FAIL】", "1. 页面显示【完成】"])
    cases = workbook["Cases"]
    cases.cell(6, 5, "PS")
    values = [cell.value for cell in cases[6]]
    values[0], values[1] = "TC-SECOND", "另一个模块"
    cases.append(values)
    values = list(values)
    values[0], values[1] = "TC-THIRD", "首页"
    cases.append(values)
    workbook.save(source)
    workbook.close()
    runtime = tmp_path / "runtime"
    workbooks, generations, runs = WorkbookRepository(runtime), GenerationRepository(runtime), RunRepository(runtime)
    imported = WorkbookService(workbooks, ExcelWorkbookImporter()).import_bytes(source.name, source.read_bytes())
    assert imported.status == WorkbookStatus.VALID, imported
    generated = GenerationService(workbooks, generations, RuleBasedFlowTranslator(), syntax_checker=PassingSyntaxChecker()).create(CreateGenerationRequest(workbook_id=imported.workbook_id, mode=RunMode.UNATTENDED))
    service = RunService(workbooks, generations, runs, GeneratedPytestPackageBuilder(runs), PreparedDevice(), ManagedPytestRunner(_settings(tmp_path)), result_writer=ResultWorkbookWriter(), background=False)
    if cancel:
        original = service.pytest_runner.run
        def run_and_cancel(run_id, package, **kwargs):
            result = original(run_id, package, **kwargs)
            service.cancel(run_id)
            return result
        service.pytest_runner.run = run_and_cancel
    result = service.create(CreateRunRequest(generation_id=generated.generation_id))
    if cancel:
        assert result.state == RunState.CANCELLED
        run = runs.get_run(result.run_id)
        assert run.case_runs[0].state == CaseRunState.FIXTURE_FAILED
        assert run.case_runs[0].cleanup_result == "unknown"
        assert run.case_runs[1].state == CaseRunState.BLOCKED
        assert result.result_workbook_available
        return
    assert result.state == RunState.COMPLETED_WITH_FAILURES, result
    run = runs.get_run(result.run_id)
    assert run.selected_case_ids == ["TC-HOME-001", "TC-THIRD", "TC-SECOND"]
    assert all(case.state == CaseRunState.FIXTURE_FAILED for case in run.case_runs)
    assert all(case.cleanup_result == "failed" for case in run.case_runs)
    assert run.case_runs[0].business_result == "passed"
    assert result.result_workbook_available
    manifest = json.loads((runs.run_dir(result.run_id) / "result/writeback-manifest.json").read_text())
    assert manifest["fixture_execution_counts"]["FX-APP-RESET"] == 1
    assert manifest["fixture_execution_counts"]["FX-LOGIN-READY"] == 2
    assert manifest["fixture_execution_counts"]["PS"] == 1
