from __future__ import annotations

import hashlib
import json
import threading
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from maestro_app.application.flow_translator import RuleBasedFlowTranslator
from maestro_app.application.generation_service import GenerationService
from maestro_app.application.ports import ManagedProcessResult, PytestPackage
from maestro_app.application.pytest_generator import GeneratedPytestPackageBuilder
from maestro_app.application.run_service import RunService
from maestro_app.application.workbook_service import WorkbookService
from maestro_app.domain.enums import CaseRunState, RunMode, RunState
from maestro_app.domain.models import ApprovalRequest, CreateGenerationRequest, CreateRunRequest
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from maestro_app.infrastructure.generation_repository import GenerationRepository
from maestro_app.infrastructure.result_workbook import ResultWorkbookWriter
from maestro_app.infrastructure.run_repository import RunRepository
from maestro_app.infrastructure.workbook_repository import WorkbookRepository
from m1_helpers import build_sample_workbook
from test_m3_repair_run import AdvancingDiagnostic, PassingSyntaxChecker, PreparedDevice


FIRST = "TC-HOME-001"
SECOND = "TC-HOME-002"


class ScriptedLifecycleRunner:
    """A complete lifecycle, including teardown, for each requested Attempt."""

    def __init__(self, first_outcomes: list[str], *, cancel_after_result=False):
        self.first_outcomes = first_outcomes
        self.cancel_after_result = cancel_after_result
        self.calls: list[tuple[str, int]] = []
        self.cancel_event: threading.Event | None = None

    def collect(self, package: PytestPackage) -> ManagedProcessResult:
        output = "\n".join(package.node_ids) + "\n"
        package.collection_path.parent.mkdir(parents=True, exist_ok=True)
        package.collection_path.write_text(output, encoding="utf-8")
        return ManagedProcessResult(0, output, package.collection_path)

    def run(self, run_id, package, *, timeout_seconds, cancel_event, on_tick=None):
        assert not cancel_event.is_set(), "must not start a cancelled Attempt"
        self.cancel_event = cancel_event
        case_id = FIRST if "tc_home_001" in package.node_ids[0] else SECOND
        ordinal = 1 + sum(item[0] == case_id for item in self.calls)
        self.calls.append((case_id, ordinal))
        outcome = self.first_outcomes[ordinal - 1] if case_id == FIRST else "passed"
        selector = {1: "确认设置", 2: r"确认设置\(1\)", 3: r"确认设置\(2\)"}[ordinal]
        output = {
            "passed": "1 passed",
            "assertion": "Assertion is false",
            "locator": f"Element not found: Text matching regex: {selector}",
            "teardown": "Fixture cleanup failed",
            "setup": "Fixture setup failed",
            "cancelled": "Execution cancelled",
            "timeout": "Execution exceeded timeout",
            "environment": "device emulator-5554 offline",
            "unknown": "unrecognized execution error",
        }[outcome]
        package.stdout_path.parent.mkdir(parents=True, exist_ok=True)
        package.stdout_path.write_text(output, encoding="utf-8")
        now = datetime.now(UTC).isoformat()
        events = []
        lifecycle = [
            ("fixture", "FX-APP-RESET", "setup", True),
            ("fixture", "FX-LOGIN-READY", "setup", outcome != "setup"),
        ]
        if outcome != "setup":
            lifecycle.append(("case", case_id, "case", outcome in {"passed", "teardown"}))
        lifecycle.append(("fixture", "FX-CLEANUP", "teardown", outcome != "teardown"))
        for kind, flow_id, phase, succeeded in lifecycle:
            base = {
                "occurred_at": now,
                "entity_type": kind,
                "entity_id": flow_id,
                "phase": phase,
                "case_id": case_id,
            }
            events.append({**base, "event": "started"})
            events.append({
                **base,
                "event": "finished",
                "succeeded": succeeded,
                "exit_code": 0 if succeeded else 1,
            })
        package.trace_path.parent.mkdir(parents=True, exist_ok=True)
        package.trace_path.write_text(
            "".join(json.dumps(item) + "\n" for item in events), encoding="utf-8"
        )
        failure = "" if outcome == "passed" else '<failure message="failed"/>'
        package.junit_path.write_text(
            f'<testsuite><testcase name="test_{case_id.lower().replace("-", "_")}">'
            f"{failure}</testcase></testsuite>", encoding="utf-8"
        )
        if on_tick:
            on_tick()
        if self.cancel_after_result:
            cancel_event.set()
        return ManagedProcessResult(
            0 if outcome == "passed" else 1,
            output,
            package.stdout_path,
            cancelled=outcome == "cancelled",
            timed_out=outcome == "timeout",
        )

    def cancel(self, run_id):
        if self.cancel_event is not None:
            self.cancel_event.set()


def run_cases(tmp_path, runner, *, continue_on_failure, diagnostic=None):
    runtime = tmp_path / "runtime"
    workbook_repository = WorkbookRepository(runtime)
    generation_repository = GenerationRepository(runtime)
    run_repository = RunRepository(runtime)
    source = build_sample_workbook(tmp_path / "failfast.xlsx")
    workbook_file = load_workbook(source)
    cases = workbook_file["Cases"]
    for column in range(1, cases.max_column + 1):
        cases.cell(7, column, cases.cell(6, column).value)
    cases["A7"] = SECOND
    cases["C7"] = "后续独立用例"
    cases["E6"] = cases["E7"] = "FX-CLEANUP"
    for column, value in enumerate([
        "FX-CLEANUP", "返回", "后置", "用例", "清理", None, 30, "是",
        "1. 点击系统返回", "1. 页面显示【首页】",
    ], start=1):
        workbook_file["Fixtures"].cell(8, column, value)
    for row in workbook_file["Config"].iter_rows(min_row=6):
        if row[0].value == "continue_on_failure":
            row[1].value = "true" if continue_on_failure else "false"
    workbook_file.save(source)
    workbook_file.close()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = WorkbookService(
        workbook_repository, ExcelWorkbookImporter()
    ).import_bytes(source.name, source.read_bytes())
    generation_service = GenerationService(
        workbook_repository, generation_repository, RuleBasedFlowTranslator(),
        syntax_checker=PassingSyntaxChecker(),
    )
    generation = generation_service.create(
        CreateGenerationRequest(workbook_id=workbook.workbook_id, mode=RunMode.REVIEW)
    )
    generation = generation_service.approve(
        generation.generation_id,
        ApprovalRequest(flow_ids=[flow.flow_id for flow in generation.flows]),
    )
    service = RunService(
        workbook_repository, generation_repository, run_repository,
        GeneratedPytestPackageBuilder(run_repository), PreparedDevice(), runner,
        diagnostic_service=diagnostic, syntax_checker=PassingSyntaxChecker(),
        result_writer=ResultWorkbookWriter(), background=False,
    )
    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    return service, run_repository, run_repository.get_run(summary.run_id)


@pytest.mark.parametrize("continue_on_failure", [False, True])
@pytest.mark.parametrize("outcome", ["assertion", "setup", "teardown", "timeout", "environment", "unknown"])
def test_final_failure_stops_only_when_configured(tmp_path, continue_on_failure, outcome):
    runner = ScriptedLifecycleRunner([outcome])
    service, repository, run = run_cases(
        tmp_path, runner, continue_on_failure=continue_on_failure
    )
    first, second = run.case_runs
    assert run.state == RunState.COMPLETED_WITH_FAILURES
    assert len(first.attempts) == 1
    assert first.state != CaseRunState.PASSED
    assert first.cleanup_result == ("failed" if outcome == "teardown" else "passed")
    assert runner.calls == ([(FIRST, 1), (SECOND, 1)] if continue_on_failure else [(FIRST, 1)])
    assert first.attempts[0].artifact_ids
    events = service.events(run.run_id)
    cleanup = [event for event in events if event.data.get("phase") == "teardown"]
    assert cleanup and cleanup[1].data["event"] == "finished"
    if continue_on_failure:
        assert second.state == CaseRunState.PASSED
        assert not any(event.type == "case.not_run" for event in events)
    else:
        assert second.state == CaseRunState.BLOCKED
        assert second.business_result == "not_run"
        assert second.cleanup_result == "not_required"
        assert not second.attempts and second.started_at is None
        skipped = next(event for event in events if event.type == "case.not_run")
        assert skipped.entity_id == SECOND
        assert skipped.data["reason"] == "continue_on_failure_false"
        assert skipped.data["blocked_by_case_id"] == FIRST
        assert skipped.sequence > cleanup[1].sequence
        report = (repository.run_dir(run.run_id) / "result/report.html").read_text()
        assert "continue_on_failure=false" in report
        result = load_workbook(repository.result_workbook_path(run.run_id), read_only=True)
        headers = {cell.value: cell.column for cell in result["Cases"][5]}
        assert result["Cases"].cell(7, headers["执行状态"]).value == "BLOCKED"
        assert result["Cases"].cell(7, headers["执行次数"]).value == 0
        result.close()


def test_repaired_case_passes_and_does_not_trigger_failfast(tmp_path):
    runner = ScriptedLifecycleRunner(["locator", "passed"])
    _, _, run = run_cases(
        tmp_path, runner, continue_on_failure=False, diagnostic=AdvancingDiagnostic()
    )
    assert run.state == RunState.COMPLETED
    assert runner.calls == [(FIRST, 1), (FIRST, 2), (SECOND, 1)]
    assert len(run.repairs) == 1
    assert all(case.state == CaseRunState.PASSED for case in run.case_runs)


def test_failfast_waits_for_all_allowed_repairs_and_keeps_evidence(tmp_path):
    runner = ScriptedLifecycleRunner(["locator", "locator", "locator"])
    service, _, run = run_cases(
        tmp_path, runner, continue_on_failure=False, diagnostic=AdvancingDiagnostic()
    )
    assert run.state == RunState.COMPLETED_WITH_FAILURES
    assert runner.calls == [(FIRST, 1), (FIRST, 2), (FIRST, 3)]
    assert len(run.repairs) == 2
    assert len(run.case_runs[0].attempts) == 3
    assert all(attempt.artifact_ids for attempt in run.case_runs[0].attempts)
    events = service.events(run.run_id)
    finishes = [event for event in events if event.type == "attempt.finished"]
    skipped = next(event for event in events if event.type == "case.not_run")
    assert skipped.sequence > finishes[-1].sequence
    assert run.case_runs[1].state == CaseRunState.BLOCKED
    assert not run.case_runs[1].attempts


@pytest.mark.parametrize("continue_on_failure", [False, True])
@pytest.mark.parametrize("outcome,cancel_after_result", [("cancelled", False), ("assertion", True)])
def test_cancellation_takes_precedence_over_continue_policy(
    tmp_path, continue_on_failure, outcome, cancel_after_result
):
    runner = ScriptedLifecycleRunner([outcome], cancel_after_result=cancel_after_result)
    service, _, run = run_cases(tmp_path, runner, continue_on_failure=continue_on_failure)
    assert run.state == RunState.CANCELLED
    assert runner.calls == [(FIRST, 1)]
    assert run.case_runs[0].cleanup_result == "passed"
    assert run.case_runs[1].state == CaseRunState.BLOCKED
    assert not run.case_runs[1].attempts
    skipped = next(event for event in service.events(run.run_id) if event.type == "case.not_run")
    assert skipped.data["reason"] == "cancelled"


def test_cancellation_during_diagnostics_prevents_next_attempt(tmp_path):
    runner = ScriptedLifecycleRunner(["locator"])

    class CancellingDiagnostic(AdvancingDiagnostic):
        def diagnose(self, **kwargs):
            result = super().diagnose(**kwargs)
            runner.cancel_event.set()
            return result

    service, _, run = run_cases(
        tmp_path, runner, continue_on_failure=False, diagnostic=CancellingDiagnostic()
    )
    assert run.state == RunState.CANCELLED
    assert runner.calls == [(FIRST, 1)]
    assert len(run.case_runs[0].attempts) == 1
    skipped = next(event for event in service.events(run.run_id) if event.type == "case.not_run")
    assert skipped.data["reason"] == "cancelled"
