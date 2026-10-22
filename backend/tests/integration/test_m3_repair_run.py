from __future__ import annotations

import hashlib
import json
import threading
from copy import copy
from datetime import UTC, datetime
from pathlib import Path

import pytest

from openpyxl import load_workbook

from maestro_app.application.flow_translator import RuleBasedFlowTranslator
from maestro_app.application.generation_service import GenerationService
from maestro_app.application.ports import ManagedProcessResult, PytestPackage
from maestro_app.application.pytest_generator import GeneratedPytestPackageBuilder
from maestro_app.application.run_service import RunService
from maestro_app.application.syntax_checker import SyntaxCheckResult
from maestro_app.application.workbook_service import WorkbookService
from maestro_app.domain.enums import CaseRunState, RunMode, RunState
from maestro_app.domain.models import CreateGenerationRequest, CreateRunRequest, DevicePreparation
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from maestro_app.infrastructure.generation_repository import GenerationRepository
from maestro_app.infrastructure.mcp_diagnostics import McpDiagnosticResult
from maestro_app.infrastructure.mcp_diagnostics import MaestroMcpDiagnosticService
from maestro_app.infrastructure.result_workbook import ResultWorkbookWriter
from maestro_app.infrastructure.run_repository import RunRepository
from maestro_app.infrastructure.workbook_repository import WorkbookRepository
from m1_helpers import build_sample_workbook
from unit.test_mcp_diagnostics import DiagnosticClient, device, screenshot


@pytest.mark.parametrize("fault", ["wrong-device", "bad-image", "valid"])
def test_real_diagnostic_gate_controls_run_repair(tmp_path, monkeypatch, fault):
    runtime = tmp_path / "runtime"
    workbooks = WorkbookRepository(runtime)
    generations = GenerationRepository(runtime)
    runs = RunRepository(runtime)
    source = build_sample_workbook(tmp_path / "diagnostic-gate.xlsx")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    imported = WorkbookService(workbooks, ExcelWorkbookImporter()).import_bytes(source.name, source.read_bytes())
    generation = GenerationService(workbooks, generations, RuleBasedFlowTranslator(), syntax_checker=PassingSyntaxChecker()).create(
        CreateGenerationRequest(workbook_id=imported.workbook_id, mode=RunMode.UNATTENDED)
    )
    client = DiagnosticClient(
        devices={"devices": [device(device_id="emulator-55540")]} if fault == "wrong-device" else None,
        image=screenshot(b"not an image") if fault == "bad-image" else None,
    )
    monkeypatch.setattr("maestro_app.infrastructure.mcp_diagnostics.MaestroMcpClient", lambda *a, **kw: client)
    runner = FailOncePytestRunner()
    service = RunService(workbooks, generations, runs, GeneratedPytestPackageBuilder(runs), PreparedDevice(), runner,
                         diagnostic_service=MaestroMcpDiagnosticService(), syntax_checker=PassingSyntaxChecker(),
                         result_writer=ResultWorkbookWriter(), background=False)
    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    run = runs.get_run(summary.run_id)
    case = service.get_case(run.run_id, "TC-HOME-001")
    yaml_path = runs.run_dir(run.run_id) / "generated/cases/TC-HOME-001.yaml"
    if fault == "valid":
        assert summary.state == RunState.COMPLETED
        assert runner.calls == 2 and len(run.repairs) == 1
        assert case.state == CaseRunState.PASSED
    else:
        assert summary.state == RunState.COMPLETED_WITH_FAILURES
        assert case.state == CaseRunState.FAILED
        assert runner.calls == 1 and len(case.attempts) == 1 and not run.repairs
        assert "tapOn: 确认设置" in yaml_path.read_text()
        assert not list(runs.run_dir(run.run_id).glob("repairs/**/selector.diff"))
        assert any(artifact.kind == "diagnostic_error" for artifact in run.artifacts)
        assert not any(event.state == RunState.REPAIRING for event in runs.list_events(run.run_id))
    assert run.result_workbook_path and (runs.run_dir(run.run_id) / run.result_workbook_path).is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert client.closed


class PassingSyntaxChecker:
    def check(self, flow_path, *, maestro_command, timeout_seconds):
        assert Path(flow_path).is_file()
        return SyntaxCheckResult(valid=True, output="OK", exit_code=0)


class PreparedDevice:
    def prepare(self, config):
        return DevicePreparation(
            device_id="emulator-5554",
            avd_name="Pixel_5",
            app_id=config.app_id,
            app_installed=True,
            boot_completed=True,
        )


class EvidenceBackedDiagnostic:
    def diagnose(self, *, maestro_command, device_id, failed_selector, evidence_dir):
        assert device_id == "emulator-5554"
        assert failed_selector == "确认设置"
        evidence_dir.mkdir(parents=True, exist_ok=True)
        hierarchy = evidence_dir / "hierarchy.json"
        hierarchy.write_text(
            json.dumps({"elements": [{"txt": "确认设置(1)"}]}, ensure_ascii=False),
            encoding="utf-8",
        )
        screenshot = evidence_dir / "screenshot.png"
        screenshot.write_bytes(b"fake-png-evidence")
        return McpDiagnosticResult(
            candidate="确认设置(1)",
            reason="页面层级中发现唯一相关文本 '确认设置(1)'",
            hierarchy_path=hierarchy,
            screenshot_path=screenshot,
            evidence_paths=(hierarchy, screenshot),
        )


class FailOncePytestRunner:
    def __init__(
        self,
        *,
        failure_selector: str = "确认设置",
        succeed_after_first: bool = True,
    ) -> None:
        self.calls = 0
        self.calls_by_case: dict[str, int] = {}
        self.failure_selector = failure_selector
        self.succeed_after_first = succeed_after_first

    def collect(self, package: PytestPackage) -> ManagedProcessResult:
        output = "\n".join(package.node_ids) + "\n"
        package.collection_path.parent.mkdir(parents=True, exist_ok=True)
        package.collection_path.write_text(output, encoding="utf-8")
        return ManagedProcessResult(0, output, package.collection_path)

    def run(
        self,
        run_id: str,
        package: PytestPackage,
        *,
        timeout_seconds: float,
        cancel_event: threading.Event,
        on_tick=None,
    ) -> ManagedProcessResult:
        self.calls += 1
        case_id = "TC-HOME-002" if "tc_home_002" in package.node_ids[0] else "TC-HOME-001"
        case_call = self.calls_by_case.get(case_id, 0) + 1
        self.calls_by_case[case_id] = case_call
        succeeded = self._succeeded(case_id, case_call)
        failure_selector = self._failure_selector(case_id, case_call)
        root = package.test_file.parents[2]
        log_path = package.stdout_path.parent / "case-flow.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        output = (
            "1 passed"
            if succeeded
            else f"Element not found: Text matching regex: {failure_selector}"
        )
        log_path.write_text(output + "\n", encoding="utf-8")
        package.stdout_path.write_text(output + "\n", encoding="utf-8")
        now = datetime.now(UTC).isoformat()
        events = []
        for fixture_id in ("FX-APP-RESET", "FX-LOGIN-READY"):
            events.extend(
                [
                    {
                        "occurred_at": now,
                        "event": "started",
                        "entity_type": "fixture",
                        "entity_id": fixture_id,
                        "phase": "setup",
                        "case_id": case_id,
                    },
                    {
                        "occurred_at": now,
                        "event": "finished",
                        "entity_type": "fixture",
                        "entity_id": fixture_id,
                        "phase": "setup",
                        "case_id": case_id,
                        "succeeded": True,
                        "exit_code": 0,
                    },
                ]
            )
        events.extend(
            [
                {
                    "occurred_at": now,
                    "event": "started",
                    "entity_type": "case",
                    "entity_id": case_id,
                    "phase": "case",
                    "case_id": case_id,
                },
                {
                    "occurred_at": now,
                    "event": "finished",
                    "entity_type": "case",
                    "entity_id": case_id,
                    "phase": "case",
                    "case_id": case_id,
                    "succeeded": succeeded,
                    "exit_code": 0 if succeeded else 1,
                    "log_path": str(log_path.relative_to(root)),
                },
            ]
        )
        package.trace_path.parent.mkdir(parents=True, exist_ok=True)
        package.trace_path.write_text(
            "".join(json.dumps(event) + "\n" for event in events), encoding="utf-8"
        )
        package.junit_path.parent.mkdir(parents=True, exist_ok=True)
        failure = "" if succeeded else "<failure message='locator failed'/>"
        package.junit_path.write_text(
            f'<testsuite tests="1" failures="{0 if succeeded else 1}">'
            f'<testcase name="test_{case_id.lower().replace("-", "_")}">{failure}</testcase></testsuite>',
            encoding="utf-8",
        )
        if on_tick:
            on_tick()
        return ManagedProcessResult(
            0 if succeeded else 1,
            output,
            package.stdout_path,
        )

    def cancel(self, run_id: str) -> None:
        return None

    def _succeeded(self, case_id: str, case_call: int) -> bool:
        return self.succeed_after_first and case_call > 1

    def _failure_selector(self, case_id: str, case_call: int) -> str:
        return self.failure_selector


class RetryCapRunner(FailOncePytestRunner):
    def _succeeded(self, case_id: str, case_call: int) -> bool:
        return case_id == "TC-HOME-002"

    def _failure_selector(self, case_id: str, case_call: int) -> str:
        return {
            1: "确认设置",
            2: r"确认设置\(1\)",
            3: r"确认设置\(2\)",
        }[case_call]


class AdvancingDiagnostic:
    def diagnose(self, *, maestro_command, device_id, failed_selector, evidence_dir):
        candidate = "确认设置(1)" if failed_selector == "确认设置" else "确认设置(2)"
        evidence_dir.mkdir(parents=True, exist_ok=True)
        hierarchy = evidence_dir / "hierarchy.json"
        hierarchy.write_text(
            json.dumps({"elements": [{"txt": candidate}]}, ensure_ascii=False),
            encoding="utf-8",
        )
        screenshot = evidence_dir / "screenshot.png"
        screenshot.write_bytes(b"fake-png-evidence")
        return McpDiagnosticResult(
            candidate=candidate,
            reason=f"唯一候选 {candidate}",
            hierarchy_path=hierarchy,
            screenshot_path=screenshot,
            evidence_paths=(hierarchy, screenshot),
        )


def test_locator_failure_is_repaired_and_retried_with_full_fixture_chain(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    workbook_repository = WorkbookRepository(runtime)
    generation_repository = GenerationRepository(runtime)
    run_repository = RunRepository(runtime)
    source = build_sample_workbook(tmp_path / "m3.xlsx")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = WorkbookService(
        workbook_repository, ExcelWorkbookImporter()
    ).import_bytes(source.name, source.read_bytes())
    generation = GenerationService(
        workbook_repository,
        generation_repository,
        RuleBasedFlowTranslator(),
        syntax_checker=PassingSyntaxChecker(),
    ).create(
        CreateGenerationRequest(
            workbook_id=workbook.workbook_id, mode=RunMode.UNATTENDED
        )
    )
    runner = FailOncePytestRunner()
    service = RunService(
        workbook_repository,
        generation_repository,
        run_repository,
        GeneratedPytestPackageBuilder(run_repository),
        PreparedDevice(),
        runner,
        diagnostic_service=EvidenceBackedDiagnostic(),
        syntax_checker=PassingSyntaxChecker(),
        result_writer=ResultWorkbookWriter(),
        background=False,
    )

    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    run = run_repository.get_run(summary.run_id)
    case = service.get_case(summary.run_id, "TC-HOME-001")

    assert summary.state == RunState.COMPLETED
    assert runner.calls == 2
    assert case.state == CaseRunState.PASSED
    assert [item.ordinal for item in case.attempts] == [1, 2]
    assert case.attempts[0].failed_step == 3
    assert case.attempts[0].repair_id == run.repairs[0].repair_id
    assert run.repairs[0].before_selector == "确认设置"
    assert run.repairs[0].after_selector == "确认设置(1)"
    assert summary.selected_case_ids == ["TC-HOME-001"]
    html_reports = [item for item in run.artifacts if item.kind == "html_report"]
    assert len(html_reports) == 1
    assert summary.html_report_artifact_id == html_reports[0].artifact_id
    report_html = (
        run_repository.run_dir(run.run_id) / html_reports[0].relative_path
    ).read_text(encoding="utf-8")
    assert "TC-HOME-001" in report_html
    assert "确认设置(1)" in report_html
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    final_yaml = (
        run_repository.run_dir(run.run_id)
        / "generated/cases/TC-HOME-001.yaml"
    ).read_text(encoding="utf-8")
    assert r"确认设置\(1\)" in final_yaml
    assert "assertVisible: 尽情发挥" in final_yaml
    for ordinal in (1, 2):
        trace = (
            run_repository.run_dir(run.run_id)
            / f"attempts/TC-HOME-001/a{ordinal}/reports/flow-trace.jsonl"
        ).read_text(encoding="utf-8")
        assert trace.count('"entity_id": "FX-APP-RESET"') == 2
        assert trace.count('"entity_id": "FX-LOGIN-READY"') == 2
    assert run.result_workbook_path is not None
    result = load_workbook(
        run_repository.run_dir(run.run_id) / run.result_workbook_path,
        read_only=True,
    )
    headers = {cell.value: cell.column for cell in result["Cases"][5]}
    assert result["Cases"].cell(6, headers["执行次数"]).value == 2
    assert result["Cases"].cell(6, headers["自动修复次数"]).value == 1
    assert result["Cases"].cell(6, headers["失败步骤"]).value == 3
    assert result["Cases"].cell(6, headers["失败类型"]).value == "LOCATOR"
    result.close()


def test_assertion_failure_is_never_repaired_or_retried(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    workbook_repository = WorkbookRepository(runtime)
    generation_repository = GenerationRepository(runtime)
    run_repository = RunRepository(runtime)
    source = build_sample_workbook(tmp_path / "m3-assertion.xlsx")
    workbook = WorkbookService(
        workbook_repository, ExcelWorkbookImporter()
    ).import_bytes(source.name, source.read_bytes())
    generation = GenerationService(
        workbook_repository,
        generation_repository,
        RuleBasedFlowTranslator(),
        syntax_checker=PassingSyntaxChecker(),
    ).create(
        CreateGenerationRequest(
            workbook_id=workbook.workbook_id, mode=RunMode.UNATTENDED
        )
    )
    runner = FailOncePytestRunner(
        failure_selector="尽情发挥",
        succeed_after_first=False,
    )
    service = RunService(
        workbook_repository,
        generation_repository,
        run_repository,
        GeneratedPytestPackageBuilder(run_repository),
        PreparedDevice(),
        runner,
        diagnostic_service=EvidenceBackedDiagnostic(),
        syntax_checker=PassingSyntaxChecker(),
        result_writer=ResultWorkbookWriter(),
        background=False,
    )

    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    run = run_repository.get_run(summary.run_id)
    case = service.get_case(summary.run_id, "TC-HOME-001")

    assert summary.state == RunState.COMPLETED_WITH_FAILURES
    assert runner.calls == 1
    assert len(case.attempts) == 1
    assert case.attempts[0].failure_type.value == "ASSERTION"
    assert case.attempts[0].failed_selector == "尽情发挥"
    assert run.repairs == []
    assert len([item for item in run.artifacts if item.kind == "html_report"]) == 1
    initial_yaml = (
        run_repository.run_dir(run.run_id)
        / "generated/initial/cases/TC-HOME-001.yaml"
    ).read_bytes()
    final_yaml = (
        run_repository.run_dir(run.run_id)
        / "generated/cases/TC-HOME-001.yaml"
    ).read_bytes()
    assert hashlib.sha256(initial_yaml).hexdigest() == hashlib.sha256(final_yaml).hexdigest()
    result = load_workbook(
        run_repository.run_dir(run.run_id) / run.result_workbook_path,
        read_only=True,
    )
    headers = {cell.value: cell.column for cell in result["Cases"][5]}
    assert result["Cases"].cell(6, headers["执行次数"]).value == 1
    assert result["Cases"].cell(6, headers["自动修复次数"]).value == 0
    assert result["Cases"].cell(6, headers["失败类型"]).value == "ASSERTION"
    result.close()


def test_retry_cap_stops_current_case_and_continues_next_case(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    workbook_repository = WorkbookRepository(runtime)
    generation_repository = GenerationRepository(runtime)
    run_repository = RunRepository(runtime)
    source = build_sample_workbook(tmp_path / "m3-cap.xlsx")
    workbook_file = load_workbook(source)
    cases = workbook_file["Cases"]
    for column in range(1, cases.max_column + 1):
        source_cell = cases.cell(6, column)
        target_cell = cases.cell(7, column)
        target_cell.value = source_cell.value
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
    cases["A7"] = "TC-HOME-002"
    cases["C7"] = "后续独立用例"
    workbook_file.save(source)
    workbook_file.close()
    workbook = WorkbookService(
        workbook_repository, ExcelWorkbookImporter()
    ).import_bytes(source.name, source.read_bytes())
    generation = GenerationService(
        workbook_repository,
        generation_repository,
        RuleBasedFlowTranslator(),
        syntax_checker=PassingSyntaxChecker(),
    ).create(
        CreateGenerationRequest(
            workbook_id=workbook.workbook_id, mode=RunMode.UNATTENDED
        )
    )
    runner = RetryCapRunner()
    service = RunService(
        workbook_repository,
        generation_repository,
        run_repository,
        GeneratedPytestPackageBuilder(run_repository),
        PreparedDevice(),
        runner,
        diagnostic_service=AdvancingDiagnostic(),
        syntax_checker=PassingSyntaxChecker(),
        result_writer=ResultWorkbookWriter(),
        background=False,
    )

    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    run = run_repository.get_run(summary.run_id)
    first = service.get_case(summary.run_id, "TC-HOME-001")
    second = service.get_case(summary.run_id, "TC-HOME-002")

    assert summary.state == RunState.COMPLETED_WITH_FAILURES
    assert [attempt.ordinal for attempt in first.attempts] == [1, 2, 3]
    assert len(first.repairs) == 2
    assert runner.calls_by_case["TC-HOME-001"] == 3
    assert second.state == CaseRunState.PASSED
    assert len(second.attempts) == 1
    assert runner.calls_by_case["TC-HOME-002"] == 1
    assert len(run.repairs) == 2
    for ordinal in (1, 2, 3):
        trace = (
            run_repository.run_dir(run.run_id)
            / f"attempts/TC-HOME-001/a{ordinal}/reports/flow-trace.jsonl"
        ).read_text(encoding="utf-8")
        assert trace.count('"entity_id": "FX-APP-RESET"') == 2
        assert trace.count('"entity_id": "FX-LOGIN-READY"') == 2
