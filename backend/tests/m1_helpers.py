from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


CASE_HEADERS = [
    "用例ID",
    "模块",
    "标题",
    "前置夹具ID",
    "后置夹具ID",
    "优先级",
    "标签",
    "是否启用",
    "操作步骤",
    "预期结果",
    "生成状态",
    "初始YAML文件",
    "执行状态",
    "执行次数",
    "自动修复次数",
    "失败步骤",
    "失败类型",
    "失败原因",
    "最终YAML文件",
    "报告路径",
    "截图/日志路径",
    "开始时间",
    "结束时间",
    "执行耗时(秒)",
    "备注",
]
FIXTURE_HEADERS = [
    "夹具ID",
    "标题",
    "阶段",
    "作用域",
    "模块",
    "依赖夹具ID",
    "执行顺序",
    "是否启用",
    "操作步骤",
    "预期结果",
    "生成状态",
    "YAML文件",
    "执行状态",
    "自动修复次数",
    "失败类型",
    "失败原因",
    "最后执行时间",
    "备注",
]
CONFIG = [
    ("project_name", "AIPPT", "是"),
    ("app_id", "aippt.com.mm", "是"),
    ("apk_path", None, "否"),
    ("android_device_id", None, "否"),
    ("avd_name", "Pixel_5", "否"),
    ("maestro_command", "maestro", "是"),
    ("execution_target", "android_emulator", "是"),
    ("output_dir", "./outputs", "是"),
    ("test_timeout_ms", 120000, "是"),
    ("max_syntax_repair", 3, "是"),
    ("max_locator_repair", 2, "是"),
    ("max_total_run_attempts", 3, "是"),
    ("continue_on_failure", "true", "是"),
    ("launch_clear_state", "false", "是"),
    ("pytest_workers", 1, "是"),
    ("result_writeback", "true", "是"),
    ("language", "zh-CN", "是"),
    ("fixture_id_separator", ";", "是"),
    ("allow_assertion_rewrite", "false", "是"),
]


def build_sample_workbook(path: Path) -> Path:
    workbook = Workbook()
    cases = workbook.active
    cases.title = "Cases"
    fixtures = workbook.create_sheet("Fixtures")
    config = workbook.create_sheet("Config")

    for column, header in enumerate(CASE_HEADERS, start=1):
        cases.cell(5, column, header)
    case_values = [
        "TC-HOME-001",
        "首页",
        "页数设置",
        "FX-LOGIN-READY",
        None,
        "P0",
        "smoke;home",
        "是",
        "1. 点击【页数】\n2. 点击【尽情发挥】\n3. 点击【确认设置】",
        "1. 【页数】当前值显示为【尽情发挥】",
    ]
    for column, value in enumerate(case_values, start=1):
        cases.cell(6, column, value)

    for column, header in enumerate(FIXTURE_HEADERS, start=1):
        fixtures.cell(5, column, header)
    fixture_rows = [
        [
            "FX-APP-RESET",
            "清除APP状态",
            "前置",
            "用例",
            "重置",
            None,
            10,
            "是",
            "1. 清除 APP 状态并启动 APP",
            "1. 页面显示【同意并接受】按钮",
        ],
        [
            "FX-LOGIN-READY",
            "登录",
            "前置",
            "用例",
            "登录",
            "FX-APP-RESET",
            20,
            "是",
            "1. 点击【同意并接受】\n2. 点击【密码登录】按钮\n3. 点击【请输入手机号/邮箱】\n4. 输入【test-user@example.invalid】\n5. 点击【请输入密码】\n6. 输入【TEST_PASSWORD_NOT_REAL】\n7. 点击【同意协议并登录】",
            "1. 页面显示【立即生成】按钮",
        ],
    ]
    for row, values in enumerate(fixture_rows, start=6):
        for column, value in enumerate(values, start=1):
            fixtures.cell(row, column, value)

    for column, value in enumerate(("配置项", "值", "必填"), start=1):
        config.cell(5, column, value)
    for row, values in enumerate(CONFIG, start=6):
        for column, value in enumerate(values, start=1):
            config.cell(row, column, value)

    workbook.save(path)
    return path
