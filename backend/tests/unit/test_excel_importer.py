import hashlib
from pathlib import Path

from openpyxl import load_workbook
import pytest

from maestro_app.domain.enums import WorkbookStatus
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from m1_helpers import build_sample_workbook


def _import(path: Path):
    return ExcelWorkbookImporter().import_workbook(
        path,
        workbook_id="wb-test",
        source_filename=path.name,
        source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )


def test_baseline_workbook_parses_all_entities_and_dependency_plan(tmp_path: Path) -> None:
    path = build_sample_workbook(tmp_path / "baseline.xlsx")

    snapshot = _import(path)

    assert snapshot.status == WorkbookStatus.VALID
    assert snapshot.counts.model_dump() == {
        "cases": 1,
        "fixtures": 2,
        "configs": 19,
        "enabled_cases": 1,
        "enabled_fixtures": 2,
        "errors": 0,
        "warnings": 0,
    }
    assert snapshot.fixture_plans["TC-HOME-001"] == [
        "FX-APP-RESET",
        "FX-LOGIN-READY",
    ]


def test_invalid_separator_and_unobservable_expectation_are_located(tmp_path: Path) -> None:
    path = build_sample_workbook(tmp_path / "invalid.xlsx")
    workbook = load_workbook(path)
    workbook["Cases"]["D6"] = "FX-LOGIN-READY；FX-APP-RESET"
    workbook["Cases"]["J6"] = "1. 页面正常"
    workbook.save(path)

    snapshot = _import(path)

    codes = {issue.code for issue in snapshot.validation_issues}
    assert snapshot.status == WorkbookStatus.INVALID
    assert {"LIST_SEPARATOR_INVALID", "EXPECTATION_NOT_OBSERVABLE"} <= codes
    expectation = next(
        issue
        for issue in snapshot.validation_issues
        if issue.code == "EXPECTATION_NOT_OBSERVABLE"
    )
    assert (expectation.sheet, expectation.row, expectation.column) == (
        "Cases",
        6,
        "J",
    )


def test_fixture_cycle_is_blocking(tmp_path: Path) -> None:
    path = build_sample_workbook(tmp_path / "cycle.xlsx")
    workbook = load_workbook(path)
    workbook["Fixtures"]["F6"] = "FX-LOGIN-READY"
    workbook.save(path)

    snapshot = _import(path)

    assert any(
        issue.code == "FIXTURE_DEPENDENCY_CYCLE"
        for issue in snapshot.validation_issues
    )


def test_required_and_enum_errors_have_exact_cells(tmp_path: Path) -> None:
    required_path = build_sample_workbook(tmp_path / "required-error.xlsx")
    workbook = load_workbook(required_path)
    workbook["Cases"]["C6"] = None
    workbook.save(required_path)
    required_snapshot = _import(required_path)

    enum_path = build_sample_workbook(tmp_path / "enum-error.xlsx")
    workbook = load_workbook(enum_path)
    workbook["Cases"]["F6"] = "P9"
    workbook.save(enum_path)
    enum_snapshot = _import(enum_path)

    required = next(
        issue
        for issue in required_snapshot.validation_issues
        if issue.code == "FIELD_REQUIRED"
    )
    enum = next(
        issue
        for issue in enum_snapshot.validation_issues
        if issue.code == "CASE_VALUE_INVALID" and issue.field == "priority"
    )
    assert (required.sheet, required.row, required.column, required.field) == (
        "Cases",
        6,
        "C",
        "title",
    )
    assert (enum.sheet, enum.row, enum.column, enum.field) == (
        "Cases",
        6,
        "F",
        "priority",
    )


def test_duplicate_id_and_config_constraint_are_blocking(tmp_path: Path) -> None:
    path = build_sample_workbook(tmp_path / "business-errors.xlsx")
    workbook = load_workbook(path)
    cases = workbook["Cases"]
    for column in range(1, cases.max_column + 1):
        cases.cell(7, column, cases.cell(6, column).value)
    config = workbook["Config"]
    pytest_workers_row = next(
        row
        for row in range(6, config.max_row + 1)
        if config.cell(row, 1).value == "pytest_workers"
    )
    config.cell(pytest_workers_row, 2, 2)
    workbook.save(path)

    snapshot = _import(path)

    duplicate = next(
        issue for issue in snapshot.validation_issues if issue.code == "ID_DUPLICATE"
    )
    conflict = next(
        issue
        for issue in snapshot.validation_issues
        if issue.code == "CONFIG_VALUE_INVALID"
        and issue.field == "pytest_workers"
    )
    assert (duplicate.sheet, duplicate.row, duplicate.column) == ("Cases", 7, "A")
    assert (conflict.sheet, conflict.row, conflict.column) == (
        "Config",
        pytest_workers_row,
        "B",
    )


def test_missing_fixture_dependency_is_located(tmp_path: Path) -> None:
    path = build_sample_workbook(tmp_path / "missing-dependency.xlsx")
    workbook = load_workbook(path)
    workbook["Fixtures"]["F7"] = "FX-NOT-FOUND"
    workbook.save(path)

    snapshot = _import(path)

    missing = next(
        issue
        for issue in snapshot.validation_issues
        if issue.code == "FIXTURE_DEPENDENCY_MISSING"
    )
    assert (missing.sheet, missing.row, missing.column) == ("Fixtures", 7, "F")


@pytest.mark.parametrize("sheet, cell, value, code", [
    ("Cases", "A5", "wrong header", "HEADER_MISSING"),
    ("Fixtures", "A5", "wrong header", "HEADER_MISSING"),
    ("Config", "A5", "wrong header", "HEADER_MISSING"),
    ("Cases", "A6", "123-invalid", "ID_FORMAT_INVALID"),
    ("Cases", "H6", "yes", "CASE_VALUE_INVALID"),
    ("Cases", "I6", "点击【按钮】", "STEP_FORMAT_INVALID"),
    ("Cases", "I6", "2. 点击【按钮】", "STEP_NUMBER_INVALID"),
    ("Cases", "I6", "1. 点击坐标 10,20", "COORDINATE_FORBIDDEN"),
    ("Cases", "I6", "1. 输入【测试】", "INPUT_WITHOUT_FOCUS"),
    ("Cases", "I6", None, "FIELD_REQUIRED"),
    ("Cases", "G6", "smoke;smoke", "LIST_VALUE_DUPLICATE"),
    ("Cases", "D6", "FX-MISSING", "CASE_FIXTURE_MISSING"),
    ("Fixtures", "H7", "否", "CASE_FIXTURE_DISABLED"),
    ("Fixtures", "C7", "后置", "CASE_FIXTURE_PHASE_MISMATCH"),
    ("Fixtures", "C7", "unknown", "FIXTURE_VALUE_INVALID"),
    ("Fixtures", "D7", "unknown", "FIXTURE_VALUE_INVALID"),
    ("Fixtures", "G7", "not an integer", "FIXTURE_VALUE_INVALID"),
    ("Fixtures", "F7", "FX-APP-RESET;FX-APP-RESET", "FIXTURE_VALUE_INVALID"),
])
def test_import_validation_boundary_matrix(tmp_path, sheet, cell, value, code):
    path = build_sample_workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)
    book.close()
    result = _import(path)
    assert result.status == WorkbookStatus.INVALID
    assert code in {issue.code for issue in result.validation_issues}


@pytest.mark.parametrize("key,value,code", [
    ("test_timeout_ms", True, "CONFIG_TYPE_INVALID"),
    ("test_timeout_ms", "invalid", "CONFIG_TYPE_INVALID"),
    ("continue_on_failure", "yes", "CONFIG_TYPE_INVALID"),
    ("app_id", None, "CONFIG_VALUE_REQUIRED"),
    ("allow_assertion_rewrite", "true", "CONFIG_VALUE_INVALID"),
])
def test_invalid_config_values_are_blocking(tmp_path, key, value, code):
    path = build_sample_workbook(tmp_path / "config.xlsx")
    book = load_workbook(path)
    for row in book["Config"].iter_rows(min_row=6):
        if row[0].value == key:
            row[1].value = value
    book.save(path)
    book.close()
    result = _import(path)
    assert result.status == WorkbookStatus.INVALID
    assert code in {issue.code for issue in result.validation_issues}


@pytest.mark.parametrize("mutation,code", [("sheet", "SHEET_MISSING"), ("key", "CONFIG_MISSING"), ("duplicate", "CONFIG_DUPLICATE")])
def test_structural_import_errors(tmp_path, mutation, code):
    path = build_sample_workbook(tmp_path / "structure.xlsx")
    book = load_workbook(path)
    if mutation == "sheet":
        del book["Config"]
    elif mutation == "key":
        book["Config"].delete_rows(6)
    else:
        book["Config"].append(["app_id", "different.app", "是"])
    book.save(path)
    book.close()
    assert code in {issue.code for issue in _import(path).validation_issues}


def test_disabled_rows_and_empty_annotation_rows_do_not_create_work(tmp_path):
    path = build_sample_workbook(tmp_path / "disabled.xlsx")
    book = load_workbook(path)
    book["Cases"]["H6"] = "否"
    book["Cases"]["Y9"] = "note only"
    book["Fixtures"]["R9"] = "note only"
    for row in book["Config"].iter_rows(min_row=6):
        if row[0].value == "continue_on_failure":
            row[1].value = False
    book.save(path)
    book.close()
    result = _import(path)
    assert result.status == WorkbookStatus.VALID
    assert len(result.cases) == 1 and len(result.fixtures) == 2
    assert result.fixture_plans == {}
    assert result.config.continue_on_failure is False


def test_unreadable_workbook_has_a_structured_error(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an xlsx")
    assert _import(path).validation_issues[0].code == "WORKBOOK_UNREADABLE"
