from __future__ import annotations

import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
import pytest

from maestro_app.domain.enums import CaseRunState, FailureType, RunMode, RunState
from maestro_app.domain.models import Artifact, Attempt, CaseRun, Run
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from maestro_app.infrastructure.result_workbook import ResultWorkbookWriter
from m1_helpers import build_sample_workbook


def test_result_writer_preserves_source_and_verifies_written_fields(tmp_path: Path) -> None:
    source = build_sample_workbook(tmp_path / "source.xlsx")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    snapshot = ExcelWorkbookImporter().import_workbook(
        source,
        workbook_id="workbook-1",
        source_filename=source.name,
        source_sha256=source_hash,
    )
    now = datetime.now(UTC)
    run = Run(
        run_id="run-result-1",
        generation_id="generation-1",
        workbook_id=snapshot.workbook_id,
        mode=RunMode.UNATTENDED,
        state=RunState.FINALIZING,
        case_runs=[
            CaseRun(
                case_id="TC-HOME-001",
                state=CaseRunState.FAILED,
                attempts=[
                    Attempt(
                        attempt_id="TC-HOME-001-a1",
                        ordinal=1,
                        state=CaseRunState.FAILED,
                        failure_type=FailureType.ASSERTION,
                        failure_reason="业务断言未满足",
                        exit_code=1,
                    )
                ],
                business_result="failed",
                cleanup_result="passed",
                started_at=now,
                ended_at=now,
                duration_seconds=1.25,
            )
        ],
    )

    evidence = ResultWorkbookWriter().write(
        source_path=source,
        result_dir=tmp_path / "result",
        snapshot=snapshot,
        run=run,
    )

    assert evidence.source_sha256_before == source_hash
    assert evidence.source_sha256_after == source_hash
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    workbook = load_workbook(evidence.path, read_only=True)
    sheet = workbook["Cases"]
    headers = {cell.value: cell.column for cell in sheet[5]}
    assert sheet.cell(6, headers["执行状态"]).value == "FAILED"
    assert sheet.cell(6, headers["执行次数"]).value == 1
    assert sheet.cell(6, headers["失败类型"]).value == "ASSERTION"
    workbook.close()


def fixture_writeback(tmp_path, events, *, stale=False):
    source = build_sample_workbook(tmp_path / "source.xlsx")
    if stale:
        book = load_workbook(source)
        for sheet_name, last_column in (("Cases", 24), ("Fixtures", 17)):
            for row in range(6, book[sheet_name].max_row + 1):
                for column in range(11, last_column + 1):
                    book[sheet_name].cell(row, column, "STALE")
        book.save(source)
        book.close()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    snapshot = ExcelWorkbookImporter().import_workbook(
        source, workbook_id="workbook-fixture", source_filename=source.name,
        source_sha256=source_hash,
    )
    root = tmp_path / "run"
    trace = root / "attempts/a1/flow-trace.jsonl"
    trace.parent.mkdir(parents=True)
    trace.write_text("\n".join(json.dumps(event) for event in events))
    generated = root / "generated/fixtures/FX-APP-RESET.yaml"
    generated.parent.mkdir(parents=True)
    generated.write_text("appId: example.app\n---\n- launchApp\n")
    run = Run(
        run_id="run-fixture", generation_id="generation-1", workbook_id=snapshot.workbook_id,
        mode=RunMode.UNATTENDED, state=RunState.FINALIZING,
        artifacts=[Artifact(
            artifact_id="trace-1", kind="trace", relative_path="attempts/a1/flow-trace.jsonl",
            media_type="application/x-ndjson", sha256=hashlib.sha256(trace.read_bytes()).hexdigest(),
        )],
    )
    return source, snapshot, root, run


def fixture_event(event, second, **extra):
    return {
        "occurred_at": f"2026-10-19T00:00:{second:02d}+00:00", "event": event,
        "entity_type": "fixture", "entity_id": "FX-APP-RESET", "phase": "setup",
        "case_id": "TC-HOME-001", **extra,
    }


@pytest.mark.parametrize("last, expected, failure", [
    ({"succeeded": True, "exit_code": 0}, "PASSED", None),
    ({"succeeded": False, "exit_code": 1}, "FAILED", "FIXTURE"),
    ({"succeeded": False, "timed_out": True, "timeout_phase": "startup"}, "FAILED", "TIMEOUT"),
    (None, "INTERRUPTED", "UNKNOWN"),
])
def test_fixture_last_actual_execution_written_and_stale_results_cleared(tmp_path, last, expected, failure):
    events = [fixture_event("started", 1), fixture_event("finished", 2, succeeded=True),
              fixture_event("started", 3)]
    if last is not None:
        events.append(fixture_event("finished", 4, **last))
    source, snapshot, root, run = fixture_writeback(tmp_path, events, stale=True)
    result = ResultWorkbookWriter().write(
        source_path=source, result_dir=root / "result", snapshot=snapshot, run=run, run_dir=root,
    )
    assert result.source_sha256_after == snapshot.source_sha256
    assert result.verified_fixture_rows == (6, 7)
    assert result.fixture_execution_counts == {"FX-APP-RESET": 2, "FX-LOGIN-READY": 0}
    book = load_workbook(result.path)
    sheet = book["Fixtures"]
    assert sheet["K6"].value == "已生成"
    assert sheet["M6"].value == expected
    assert sheet["N6"].value == 0
    assert sheet["O6"].value == failure
    assert sheet["Q6"].value.second == (4 if last is not None else 3)
    assert sheet["K7"].value == "未生成"
    assert sheet["M7"].value == "NOT_RUN"
    assert sheet["L7"].value is None
    assert sheet["P7"].value is None
    assert book["Cases"]["M6"].value == "NOT_RUN"
    assert book["Cases"]["R6"].value is None
    book.close()


@pytest.mark.parametrize("mutation", ["changed", "missing", "escape", "bad_json", "no_timezone"])
def test_fixture_corrupt_or_unsafe_evidence_blocks_writeback(tmp_path, mutation):
    source, snapshot, root, run = fixture_writeback(tmp_path, [fixture_event("started", 1)])
    artifact = run.artifacts[0]
    trace = root / artifact.relative_path
    if mutation == "changed":
        trace.write_text("changed")
    elif mutation == "missing":
        trace.unlink()
    elif mutation == "escape":
        run.artifacts[0] = artifact.model_copy(update={"relative_path": "../outside.jsonl"})
    else:
        trace.write_text("not json" if mutation == "bad_json" else json.dumps(fixture_event("started", 1, occurred_at="2026-10-19T00:00:01")))
        run.artifacts[0] = artifact.model_copy(update={"sha256": hashlib.sha256(trace.read_bytes()).hexdigest()})
    with pytest.raises(RuntimeError):
        ResultWorkbookWriter().write(source_path=source, result_dir=root / "result", snapshot=snapshot, run=run)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == snapshot.source_sha256
    assert not (root / "result").exists()


def test_fixture_failure_reason_is_literal_text_not_formula(tmp_path):
    source, snapshot, root, run = fixture_writeback(tmp_path, [
        fixture_event("started", 1), fixture_event("finished", 2, succeeded=False, failure_reason='=HYPERLINK("https://example.invalid")'),
    ])
    result = ResultWorkbookWriter().write(source_path=source, result_dir=root / "result", snapshot=snapshot, run=run)
    book = load_workbook(result.path)
    assert book["Fixtures"]["P6"].data_type == "s"
    book.close()
