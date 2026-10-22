from __future__ import annotations

import hashlib
import json
import os
import tempfile
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from maestro_app.domain.models import Artifact, Run, WorkbookSnapshot


@dataclass(frozen=True, slots=True)
class ResultWorkbookEvidence:
    path: Path
    source_sha256_before: str
    source_sha256_after: str
    output_sha256: str
    verified_rows: tuple[int, ...]
    verified_fixture_rows: tuple[int, ...] = ()
    fixture_execution_counts: dict[str, int] = field(default_factory=dict)


class ResultWorkbookWriter:
    CASE_RESULT_FIELDS = (
        "生成状态",
        "初始YAML文件",
        "执行状态",
        "执行次数",
        "自动修复次数",
        "失败步骤",
        "失败类型",
        "失败原因",
        "最终YAML文件",
        "报告路径",
        "截图/日志路径",
        "开始时间",
        "结束时间",
        "执行耗时(秒)",
    )
    FIXTURE_RESULT_FIELDS = (
        "生成状态",
        "YAML文件",
        "执行状态",
        "自动修复次数",
        "失败类型",
        "失败原因",
        "最后执行时间",
    )

    def write(
        self,
        *,
        source_path: Path,
        result_dir: Path,
        snapshot: WorkbookSnapshot,
        run: Run,
        run_dir: Path | None = None,
    ) -> ResultWorkbookEvidence:
        before = self._sha(source_path)
        if before != snapshot.source_sha256:
            raise RuntimeError("source workbook hash no longer matches the snapshot")
        root = (run_dir or result_dir.parent).resolve()
        observations = self._fixture_observations(run, root)
        workbook = load_workbook(source_path)
        # Every expected result cell is verified after serialization, not just status.
        expected: dict[str, dict[tuple[int, int], object]] = {
            "Cases": {}, "Fixtures": {},
        }
        verified_rows: list[int] = []
        verified_fixture_rows: list[int] = []
        fixture_counts: dict[str, int] = {}
        try:
            cases_sheet = workbook["Cases"]
            fixtures_sheet = workbook["Fixtures"]
            headers = self._headers(cases_sheet, self.CASE_RESULT_FIELDS)
            fixture_headers = self._headers(fixtures_sheet, self.FIXTURE_RESULT_FIELDS)
            case_runs = {case.case_id: case for case in run.case_runs}
            for source in snapshot.cases:
                self._verify_source_id(cases_sheet, source.source_row, source.case_id)
                case_run = case_runs.get(source.case_id)
                values = self._case_values(source, case_run, run, root)
                self._write_values(cases_sheet, source.source_row, headers, values, expected["Cases"])
                verified_rows.append(source.source_row)

            for fixture in snapshot.fixtures:
                self._verify_source_id(fixtures_sheet, fixture.source_row, fixture.fixture_id)
                executions = observations.get(fixture.fixture_id, [])
                fixture_counts[fixture.fixture_id] = len(executions)
                values = self._fixture_values(fixture, executions, run, root)
                self._write_values(fixtures_sheet, fixture.source_row, fixture_headers, values, expected["Fixtures"])
                verified_fixture_rows.append(fixture.source_row)

            result_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%S%fZ")
            output = result_dir / f"{source_path.stem}_结果_{stamp}.xlsx"
            handle, temp_name = tempfile.mkstemp(
                prefix=f".{output.name}.", suffix=".tmp.xlsx", dir=result_dir
            )
            os.close(handle)
            temp_path = Path(temp_name)
            try:
                workbook.save(temp_path)
                check = load_workbook(temp_path, data_only=False, read_only=True)
                try:
                    for sheet_name, cells in expected.items():
                        for (row, column), value in cells.items():
                            if check[sheet_name].cell(row, column).value != value:
                                raise RuntimeError(
                                    f"result workbook verification failed at {sheet_name}!{row}:{column}"
                                )
                finally:
                    check.close()
                if before != self._sha(source_path):
                    raise RuntimeError("source workbook was modified during result writeback")
                os.replace(temp_path, output)
            finally:
                if temp_path.exists():
                    temp_path.unlink()
        finally:
            workbook.close()
        return ResultWorkbookEvidence(
            path=output,
            source_sha256_before=before,
            source_sha256_after=self._sha(source_path),
            output_sha256=self._sha(output),
            verified_rows=tuple(verified_rows),
            verified_fixture_rows=tuple(verified_fixture_rows),
            fixture_execution_counts=fixture_counts,
        )

    def _case_values(self, source, case_run, run: Run, root: Path) -> dict:
        initial = self._generated_path(root, f"generated/initial/cases/{source.case_id}.yaml")
        final = self._generated_path(root, f"generated/cases/{source.case_id}.yaml")
        values = dict.fromkeys(self.CASE_RESULT_FIELDS)
        values.update({
            "生成状态": "已生成" if final else "未生成",
            "初始YAML文件": initial,
            "最终YAML文件": final,
            "执行状态": "NOT_RUN" if source.enabled else "DISABLED",
            "执行次数": 0,
            "自动修复次数": 0,
        })
        if case_run is not None:
            latest_failure = next(
                (
                    item
                    for item in reversed(case_run.attempts)
                    if item.failure_type is not None
                ),
                None,
            )
            repairs = [item for item in run.repairs if item.case_id == case_run.case_id]
            evidence = [
                item.relative_path
                for item in run.artifacts
                if item.case_id == case_run.case_id
                and item.kind in {"screenshot", "log", "hierarchy"}
            ]
            report = next(
                (
                    item.relative_path
                    for item in reversed(run.artifacts)
                    if item.case_id == case_run.case_id
                    and item.kind in {"pytest_junit", "maestro_junit"}
                ),
                None,
            )
            values.update({
                "执行状态": case_run.state.value,
                "执行次数": len(case_run.attempts),
                "自动修复次数": len(repairs),
                "失败步骤": latest_failure.failed_step if latest_failure else None,
                "失败类型": (
                    latest_failure.failure_type.value
                    if latest_failure and latest_failure.failure_type
                    else None
                ),
                "失败原因": latest_failure.failure_reason if latest_failure else None,
                "报告路径": report,
                "截图/日志路径": "\n".join(evidence) or None,
                "开始时间": self._excel_datetime(case_run.started_at),
                "结束时间": self._excel_datetime(case_run.ended_at),
                "执行耗时(秒)": case_run.duration_seconds,
            })
        return values

    def _fixture_values(self, fixture, executions: list[dict], run: Run, root: Path) -> dict:
        yaml_path = self._generated_path(root, f"generated/fixtures/{fixture.fixture_id}.yaml")
        values = dict.fromkeys(self.FIXTURE_RESULT_FIELDS)
        values.update({
            "生成状态": "已生成" if yaml_path else "未生成",
            "YAML文件": yaml_path,
            "执行状态": "NOT_RUN" if fixture.enabled else "DISABLED",
            # Fixture failures are not eligible for the Case locator-repair policy.
            "自动修复次数": 0,
        })
        if not executions:
            return values
        last = executions[-1]
        values["最后执行时间"] = self._excel_datetime(last["time"])
        if last.get("event") == "finished" and last.get("succeeded") is True:
            values["执行状态"] = "PASSED"
        elif last.get("event") == "finished" and last.get("succeeded") is False:
            values["执行状态"] = "FAILED"
            values["失败类型"] = "TIMEOUT" if last.get("timed_out") else "FIXTURE"
            detail = last.get("failure_reason") or last.get("reason")
            if not detail:
                detail = (
                    f"{last.get('timeout_phase') or 'execution'} timed out"
                    if last.get("timed_out")
                    else f"Maestro fixture exited with code {last.get('exit_code', 'unknown')}"
                )
            values["失败原因"] = str(detail)
        else:
            values["执行状态"] = "INTERRUPTED"
            values["失败类型"] = "UNKNOWN"
            values["失败原因"] = "夹具已开始但没有完整结束证据" + (
                f"；{run.terminal_reason}" if run.terminal_reason else ""
            )
        return values

    def _fixture_observations(self, run: Run, root: Path) -> dict[str, list[dict]]:
        events: list[tuple[datetime, int, str, dict]] = []
        seen: set[str] = set()
        for artifact in run.artifacts:
            if artifact.kind != "trace" or artifact.relative_path in seen:
                continue
            seen.add(artifact.relative_path)
            path = self._artifact_path(root, artifact)
            for line in path.read_text(encoding="utf-8").splitlines():
                if not line.strip():
                    continue
                try:
                    event = json.loads(line)
                    if not isinstance(event, dict):
                        raise ValueError("trace event must be an object")
                    if event.get("entity_type") != "fixture" or event.get("event") not in {"started", "finished"}:
                        continue
                    occurred = datetime.fromisoformat(event["occurred_at"])
                    if occurred.tzinfo is None or not event.get("entity_id"):
                        raise ValueError("fixture event must have an ID and timezone")
                    events.append((occurred, len(events), artifact.relative_path, event))
                except (KeyError, TypeError, ValueError) as exc:
                    raise RuntimeError(f"invalid registered fixture trace: {artifact.relative_path}") from exc
        observations: dict[str, list[dict]] = {}
        pending: dict[tuple[str, str, str, str], dict] = {}
        for occurred, _, trace, event in sorted(events, key=lambda item: (item[0], item[1])):
            fixture_id = event["entity_id"]
            key = (trace, fixture_id, event.get("case_id", ""), event.get("phase", ""))
            value = {**event, "time": occurred}
            if event["event"] == "finished" and key in pending:
                pending.pop(key).update(value)
            else:
                observations.setdefault(fixture_id, []).append(value)
                if event["event"] == "started":
                    pending[key] = value
        return observations

    @classmethod
    def _artifact_path(cls, root: Path, artifact: Artifact) -> Path:
        path = cls._safe_path(root, artifact.relative_path)
        if not path.is_file() or cls._sha(path) != artifact.sha256:
            raise RuntimeError(f"registered trace is missing or changed: {artifact.relative_path}")
        return path

    @staticmethod
    def _safe_path(root: Path, relative: str) -> Path:
        fragment = Path(relative)
        if not relative or fragment.is_absolute() or ".." in fragment.parts or "\\" in relative:
            raise RuntimeError("artifact path must be a safe Run-relative path")
        path = (root / fragment).resolve()
        if not path.is_relative_to(root):
            raise RuntimeError("artifact path escapes its Run directory")
        return path

    @classmethod
    def _generated_path(cls, root: Path, relative: str) -> str | None:
        return relative if cls._safe_path(root, relative).is_file() else None

    @staticmethod
    def _headers(sheet, fields: tuple[str, ...]) -> dict[str, int]:
        headers = {str(cell.value).strip(): cell.column for cell in sheet[5] if cell.value is not None}
        missing = [name for name in fields if name not in headers]
        if missing:
            raise RuntimeError(f"{sheet.title} result columns are missing: {missing}")
        return headers

    @staticmethod
    def _verify_source_id(sheet, row: int, entity_id: str) -> None:
        if str(sheet.cell(row, 1).value).strip() != entity_id:
            raise RuntimeError(f"source ID does not match snapshot at {sheet.title}!{row}")

    @staticmethod
    def _write_values(sheet, row: int, headers: dict[str, int], values: dict, expected: dict) -> None:
        for name, value in values.items():
            cell = sheet.cell(row, headers[name])
            # .cell(..., value=None) does not clear an existing value in openpyxl.
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            expected[(row, headers[name])] = value

    @staticmethod
    def _sha(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    @staticmethod
    def _excel_datetime(value: datetime | None) -> datetime | None:
        if value is None:
            return None
        # Excel serializes datetimes to millisecond precision.
        utc = value.astimezone(UTC).replace(tzinfo=None)
        return utc.replace(microsecond=(utc.microsecond // 1000) * 1000)
