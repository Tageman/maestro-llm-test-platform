from __future__ import annotations

import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from maestro_app.domain.enums import (
    FixturePhase,
    FixtureScope,
    Severity,
    WorkbookStatus,
)
from maestro_app.domain.fixture_graph import FixtureDependencyResolver
from maestro_app.domain.models import (
    FixtureDefinition,
    ProjectConfig,
    SourceStep,
    TestCaseDefinition,
    ValidationIssue,
    WorkbookSnapshot,
)


REQUIRED_SHEETS = ("Cases", "Fixtures", "Config")
HEADER_ROW = 5

CASE_HEADERS = {
    "case_id": "用例ID",
    "module": "模块",
    "title": "标题",
    "pre_fixture_ids": "前置夹具ID",
    "post_fixture_ids": "后置夹具ID",
    "priority": "优先级",
    "tags": "标签",
    "enabled": "是否启用",
    "steps": "操作步骤",
    "expected_results": "预期结果",
    "note": "备注",
}
FIXTURE_HEADERS = {
    "fixture_id": "夹具ID",
    "title": "标题",
    "phase": "阶段",
    "scope": "作用域",
    "module": "模块",
    "dependency_ids": "依赖夹具ID",
    "order": "执行顺序",
    "enabled": "是否启用",
    "steps": "操作步骤",
    "expected_results": "预期结果",
    "note": "备注",
}

CONFIG_FIELDS = {
    "project_name",
    "app_id",
    "apk_path",
    "android_device_id",
    "avd_name",
    "maestro_command",
    "execution_target",
    "output_dir",
    "test_timeout_ms",
    "max_syntax_repair",
    "max_locator_repair",
    "max_total_run_attempts",
    "continue_on_failure",
    "launch_clear_state",
    "pytest_workers",
    "result_writeback",
    "language",
    "fixture_id_separator",
    "allow_assertion_rewrite",
}
INTEGER_CONFIGS = {
    "test_timeout_ms",
    "max_syntax_repair",
    "max_locator_repair",
    "max_total_run_attempts",
    "pytest_workers",
}
BOOLEAN_CONFIGS = {
    "continue_on_failure",
    "launch_clear_state",
    "result_writeback",
    "allow_assertion_rewrite",
}
OPTIONAL_CONFIGS = {"apk_path", "android_device_id", "avd_name"}


class ExcelWorkbookImporter:
    def import_workbook(
        self,
        source_path: Path,
        *,
        workbook_id: str,
        source_filename: str,
        source_sha256: str,
    ) -> WorkbookSnapshot:
        issues: list[ValidationIssue] = []
        try:
            workbook = load_workbook(source_path, data_only=False, read_only=False)
        except Exception as exc:
            return WorkbookSnapshot(
                workbook_id=workbook_id,
                source_filename=source_filename,
                source_sha256=source_sha256,
                imported_at=datetime.now(UTC),
                status=WorkbookStatus.INVALID,
                validation_issues=[
                    self._issue(
                        "WORKBOOK_UNREADABLE",
                        Severity.ERROR,
                        f"无法读取工作簿：{exc}",
                        suggestion="确认文件为有效的 .xlsx 工作簿",
                    )
                ],
            )

        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                issues.append(
                    self._issue(
                        "SHEET_MISSING",
                        Severity.ERROR,
                        f"缺少必需 Sheet：{sheet_name}",
                        sheet=sheet_name,
                        suggestion="使用模板并保留 Cases、Fixtures、Config 三个 Sheet 名称",
                    )
                )
        if any(issue.code == "SHEET_MISSING" for issue in issues):
            return self._snapshot(
                workbook_id,
                source_filename,
                source_sha256,
                [],
                [],
                None,
                {},
                issues,
                {},
            )

        raw_config, config, config_issues = self._parse_config(workbook["Config"])
        issues.extend(config_issues)
        separator = (
            config.fixture_id_separator
            if config is not None
            else str(raw_config.get("fixture_id_separator") or ";")
        )
        cases, case_issues = self._parse_cases(workbook["Cases"], separator)
        fixtures, fixture_issues = self._parse_fixtures(
            workbook["Fixtures"], separator
        )
        issues.extend(case_issues)
        issues.extend(fixture_issues)
        issues.extend(self._validate_ids(cases, fixtures))

        resolver = FixtureDependencyResolver(fixtures)
        issues.extend(resolver.validate())
        reference_issues, fixture_plans = self._validate_case_references(
            cases, fixtures, resolver
        )
        issues.extend(reference_issues)

        return self._snapshot(
            workbook_id,
            source_filename,
            source_sha256,
            cases,
            fixtures,
            config,
            raw_config,
            issues,
            fixture_plans,
        )

    def _parse_config(
        self, sheet
    ) -> tuple[dict[str, str | int | bool | None], ProjectConfig | None, list[ValidationIssue]]:
        issues: list[ValidationIssue] = []
        header = [self._text(cell.value) for cell in sheet[HEADER_ROW]]
        expected = ["配置项", "值", "必填"]
        for index, value in enumerate(expected, start=1):
            if index > len(header) or header[index - 1] != value:
                issues.append(
                    self._issue(
                        "HEADER_MISSING",
                        Severity.ERROR,
                        f"Config 第 {get_column_letter(index)} 列应为“{value}”",
                        sheet="Config",
                        row=HEADER_ROW,
                        column=get_column_letter(index),
                    )
                )

        raw: dict[str, str | int | bool | None] = {}
        rows: dict[str, int] = {}
        required_flags: dict[str, bool] = {}
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            key = self._text(sheet.cell(row, 1).value)
            if not key or key not in CONFIG_FIELDS:
                continue
            if key in raw:
                issues.append(
                    self._issue(
                        "CONFIG_DUPLICATE",
                        Severity.ERROR,
                        f"配置项 {key} 重复",
                        sheet="Config",
                        row=row,
                        column="A",
                        entity_id=key,
                        field=key,
                    )
                )
                continue
            raw[key] = sheet.cell(row, 2).value
            rows[key] = row
            required_flags[key] = self._text(sheet.cell(row, 3).value) == "是"

        for key in sorted(CONFIG_FIELDS):
            if key not in raw:
                issues.append(
                    self._issue(
                        "CONFIG_MISSING",
                        Severity.ERROR,
                        f"缺少配置项 {key}",
                        sheet="Config",
                        column="A",
                        entity_id=key,
                        field=key,
                    )
                )
                continue
            if required_flags.get(key) and self._is_blank(raw[key]):
                issues.append(
                    self._issue(
                        "CONFIG_VALUE_REQUIRED",
                        Severity.ERROR,
                        f"配置项 {key} 的值不能为空",
                        sheet="Config",
                        row=rows[key],
                        column="B",
                        entity_id=key,
                        field=key,
                    )
                )

        converted: dict[str, Any] = {}
        for key, value in raw.items():
            if self._is_blank(value):
                converted[key] = None if key in OPTIONAL_CONFIGS else value
                continue
            try:
                if key in INTEGER_CONFIGS:
                    if isinstance(value, bool):
                        raise ValueError("布尔值不是整数")
                    converted[key] = int(value)
                elif key in BOOLEAN_CONFIGS:
                    converted[key] = self._parse_config_bool(value)
                else:
                    converted[key] = self._text(value)
            except (TypeError, ValueError) as exc:
                issues.append(
                    self._issue(
                        "CONFIG_TYPE_INVALID",
                        Severity.ERROR,
                        f"配置项 {key} 类型错误：{exc}",
                        sheet="Config",
                        row=rows[key],
                        column="B",
                        entity_id=key,
                        field=key,
                    )
                )

        config: ProjectConfig | None = None
        if CONFIG_FIELDS.issubset(converted):
            try:
                config = ProjectConfig.model_validate(converted)
            except ValidationError as exc:
                for error in exc.errors():
                    field = str(error["loc"][0])
                    issues.append(
                        self._issue(
                            "CONFIG_VALUE_INVALID",
                            Severity.ERROR,
                            f"配置项 {field} 不合法：{error['msg']}",
                            sheet="Config",
                            row=rows.get(field),
                            column="B",
                            entity_id=field,
                            field=field,
                        )
                    )
        return raw, config, issues

    def _parse_cases(
        self, sheet, separator: str
    ) -> tuple[list[TestCaseDefinition], list[ValidationIssue]]:
        headers, issues = self._headers(sheet, "Cases", CASE_HEADERS)
        cases: list[TestCaseDefinition] = []
        if issues:
            return cases, issues
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            raw = {
                field: sheet.cell(row, headers[field]).value for field in CASE_HEADERS
            }
            if all(self._is_blank(raw[field]) for field in CASE_HEADERS if field != "note"):
                continue
            row_issues: list[ValidationIssue] = []
            required = (
                "case_id",
                "module",
                "title",
                "priority",
                "enabled",
                "steps",
                "expected_results",
            )
            self._required_fields(
                raw, required, "Cases", row, headers, CASE_HEADERS, row_issues
            )
            case_id = self._text(raw["case_id"])
            steps = self._parse_steps(
                raw["steps"], "Cases", row, headers["steps"], case_id, "steps", row_issues
            )
            expected = self._parse_steps(
                raw["expected_results"],
                "Cases",
                row,
                headers["expected_results"],
                case_id,
                "expected_results",
                row_issues,
            )
            self._validate_step_language(
                steps, expected, "Cases", row, case_id, headers, row_issues
            )
            try:
                enabled = self._parse_enabled(raw["enabled"])
                pre = self._split_list(
                    raw["pre_fixture_ids"], separator, "Cases", row, headers["pre_fixture_ids"], case_id, "pre_fixture_ids", row_issues
                )
                post = self._split_list(
                    raw["post_fixture_ids"], separator, "Cases", row, headers["post_fixture_ids"], case_id, "post_fixture_ids", row_issues
                )
                tags = self._split_list(
                    raw["tags"], ";", "Cases", row, headers["tags"], case_id, "tags", row_issues
                )
            except ValueError as exc:
                row_issues.append(
                    self._issue(
                        "CASE_VALUE_INVALID",
                        Severity.ERROR,
                        str(exc),
                        sheet="Cases",
                        row=row,
                        entity_id=case_id or None,
                    )
                )
                enabled, pre, post, tags = False, [], [], []

            if not row_issues:
                try:
                    cases.append(
                        TestCaseDefinition(
                            case_id=case_id,
                            module=self._text(raw["module"]),
                            title=self._text(raw["title"]),
                            pre_fixture_ids=pre,
                            post_fixture_ids=post,
                            priority=self._text(raw["priority"]),
                            tags=tags,
                            enabled=enabled,
                            steps=steps,
                            expected_results=expected,
                            source_row=row,
                            note=self._text(raw["note"]) or None,
                        )
                    )
                except ValidationError as exc:
                    row_issues.extend(
                        self._pydantic_issues(
                            exc, "CASE_VALUE_INVALID", "Cases", row, case_id, headers
                        )
                    )
            issues.extend(row_issues)
        return cases, issues

    def _parse_fixtures(
        self, sheet, separator: str
    ) -> tuple[list[FixtureDefinition], list[ValidationIssue]]:
        headers, issues = self._headers(sheet, "Fixtures", FIXTURE_HEADERS)
        fixtures: list[FixtureDefinition] = []
        if issues:
            return fixtures, issues
        for row in range(HEADER_ROW + 1, sheet.max_row + 1):
            raw = {
                field: sheet.cell(row, headers[field]).value for field in FIXTURE_HEADERS
            }
            if all(self._is_blank(raw[field]) for field in FIXTURE_HEADERS if field != "note"):
                continue
            row_issues: list[ValidationIssue] = []
            required = (
                "fixture_id",
                "title",
                "phase",
                "scope",
                "order",
                "enabled",
                "steps",
                "expected_results",
            )
            self._required_fields(
                raw, required, "Fixtures", row, headers, FIXTURE_HEADERS, row_issues
            )
            fixture_id = self._text(raw["fixture_id"])
            steps = self._parse_steps(
                raw["steps"], "Fixtures", row, headers["steps"], fixture_id, "steps", row_issues
            )
            expected = self._parse_steps(
                raw["expected_results"],
                "Fixtures",
                row,
                headers["expected_results"],
                fixture_id,
                "expected_results",
                row_issues,
            )
            self._validate_step_language(
                steps, expected, "Fixtures", row, fixture_id, headers, row_issues
            )
            try:
                dependencies = self._split_list(
                    raw["dependency_ids"], separator, "Fixtures", row, headers["dependency_ids"], fixture_id, "dependency_ids", row_issues
                )
                fixture = FixtureDefinition(
                    fixture_id=fixture_id,
                    title=self._text(raw["title"]),
                    phase=self._parse_phase(raw["phase"]),
                    scope=self._parse_scope(raw["scope"]),
                    module=self._text(raw["module"]) or None,
                    dependency_ids=dependencies,
                    order=int(raw["order"]),
                    enabled=self._parse_enabled(raw["enabled"]),
                    steps=steps,
                    expected_results=expected,
                    source_row=row,
                    note=self._text(raw["note"]) or None,
                )
                if not row_issues:
                    fixtures.append(fixture)
            except (TypeError, ValueError, ValidationError) as exc:
                if isinstance(exc, ValidationError):
                    row_issues.extend(
                        self._pydantic_issues(
                            exc,
                            "FIXTURE_VALUE_INVALID",
                            "Fixtures",
                            row,
                            fixture_id,
                            headers,
                        )
                    )
                else:
                    row_issues.append(
                        self._issue(
                            "FIXTURE_VALUE_INVALID",
                            Severity.ERROR,
                            str(exc),
                            sheet="Fixtures",
                            row=row,
                            entity_id=fixture_id or None,
                        )
                    )
            issues.extend(row_issues)
        return fixtures, issues

    def _validate_ids(
        self,
        cases: list[TestCaseDefinition],
        fixtures: list[FixtureDefinition],
    ) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        seen: dict[str, tuple[str, int]] = {}
        id_pattern = re.compile(r"^[A-Za-z][A-Za-z0-9_-]*$")
        for entity, sheet, row, column in [
            *((case, "Cases", case.source_row, "A") for case in cases),
            *((fixture, "Fixtures", fixture.source_row, "A") for fixture in fixtures),
        ]:
            entity_id = getattr(entity, "case_id", None) or getattr(
                entity, "fixture_id"
            )
            if not id_pattern.fullmatch(entity_id):
                issues.append(
                    self._issue(
                        "ID_FORMAT_INVALID",
                        Severity.ERROR,
                        f"ID {entity_id} 只能包含字母、数字、下划线和连字符，且必须以字母开头",
                        sheet=sheet,
                        row=row,
                        column=column,
                        entity_id=entity_id,
                        field="id",
                    )
                )
            if entity_id in seen:
                previous_sheet, previous_row = seen[entity_id]
                issues.append(
                    self._issue(
                        "ID_DUPLICATE",
                        Severity.ERROR,
                        f"ID {entity_id} 与 {previous_sheet}!A{previous_row} 重复",
                        sheet=sheet,
                        row=row,
                        column=column,
                        entity_id=entity_id,
                        field="id",
                    )
                )
            else:
                seen[entity_id] = (sheet, row)
        return issues

    def _validate_case_references(
        self,
        cases: list[TestCaseDefinition],
        fixtures: list[FixtureDefinition],
        resolver: FixtureDependencyResolver,
    ) -> tuple[list[ValidationIssue], dict[str, list[str]]]:
        issues: list[ValidationIssue] = []
        fixture_map = {fixture.fixture_id: fixture for fixture in fixtures}
        plans: dict[str, list[str]] = {}
        for case in cases:
            if not case.enabled:
                continue
            for field, fixture_ids, expected_phase, column in (
                ("pre_fixture_ids", case.pre_fixture_ids, FixturePhase.SETUP, "D"),
                ("post_fixture_ids", case.post_fixture_ids, FixturePhase.TEARDOWN, "E"),
            ):
                for fixture_id in fixture_ids:
                    fixture = fixture_map.get(fixture_id)
                    if fixture is None:
                        issues.append(
                            self._case_ref_issue(
                                "CASE_FIXTURE_MISSING",
                                case,
                                field,
                                column,
                                f"引用夹具 {fixture_id} 不存在",
                            )
                        )
                    elif not fixture.enabled:
                        issues.append(
                            self._case_ref_issue(
                                "CASE_FIXTURE_DISABLED",
                                case,
                                field,
                                column,
                                f"启用用例不能引用停用夹具 {fixture_id}",
                            )
                        )
                    elif fixture.phase != expected_phase:
                        issues.append(
                            self._case_ref_issue(
                                "CASE_FIXTURE_PHASE_MISMATCH",
                                case,
                                field,
                                column,
                                f"夹具 {fixture_id} 阶段与引用位置不一致",
                            )
                        )

            if not any(issue.entity_id == case.case_id for issue in issues):
                try:
                    setup = resolver.resolve(
                        case.pre_fixture_ids, phase=FixturePhase.SETUP
                    )
                    teardown = resolver.resolve(
                        case.post_fixture_ids, phase=FixturePhase.TEARDOWN
                    )
                except ValueError:
                    continue
                plans[case.case_id] = setup + list(reversed(teardown))
        return issues, plans

    def _headers(
        self, sheet, sheet_name: str, expected: dict[str, str]
    ) -> tuple[dict[str, int], list[ValidationIssue]]:
        values = {
            self._text(cell.value): cell.column
            for cell in sheet[HEADER_ROW]
            if self._text(cell.value)
        }
        headers: dict[str, int] = {}
        issues: list[ValidationIssue] = []
        for field, label in expected.items():
            column = values.get(label)
            if column is None:
                issues.append(
                    self._issue(
                        "HEADER_MISSING",
                        Severity.ERROR,
                        f"{sheet_name} 缺少列“{label}”",
                        sheet=sheet_name,
                        row=HEADER_ROW,
                        field=field,
                    )
                )
            else:
                headers[field] = column
        return headers, issues

    def _required_fields(
        self,
        raw: dict[str, Any],
        required: tuple[str, ...],
        sheet: str,
        row: int,
        headers: dict[str, int],
        labels: dict[str, str],
        issues: list[ValidationIssue],
    ) -> None:
        entity_id = self._text(raw.get("case_id") or raw.get("fixture_id")) or None
        for field in required:
            if self._is_blank(raw[field]):
                issues.append(
                    self._issue(
                        "FIELD_REQUIRED",
                        Severity.ERROR,
                        f"{labels[field]}不能为空",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(headers[field]),
                        entity_id=entity_id,
                        field=field,
                    )
                )

    def _parse_steps(
        self,
        value: Any,
        sheet: str,
        row: int,
        column: int,
        entity_id: str,
        field: str,
        issues: list[ValidationIssue],
    ) -> list[SourceStep]:
        if self._is_blank(value):
            return []
        lines = [line.strip() for line in str(value).splitlines() if line.strip()]
        steps: list[SourceStep] = []
        for expected_ordinal, line in enumerate(lines, start=1):
            match = re.fullmatch(r"\s*(\d+)[.、]\s*(.+?)\s*", line)
            if match is None:
                issues.append(
                    self._issue(
                        "STEP_FORMAT_INVALID",
                        Severity.ERROR,
                        f"步骤必须逐行编号：{line}",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(column),
                        entity_id=entity_id or None,
                        field=field,
                        suggestion="使用“1. 步骤内容”的格式，每行一条",
                    )
                )
                continue
            ordinal = int(match.group(1))
            if ordinal != expected_ordinal:
                issues.append(
                    self._issue(
                        "STEP_NUMBER_INVALID",
                        Severity.ERROR,
                        f"步骤编号应连续，期望 {expected_ordinal}，实际 {ordinal}",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(column),
                        entity_id=entity_id or None,
                        field=field,
                    )
                )
            steps.append(SourceStep(ordinal=ordinal, text=match.group(2)))
        return steps

    def _validate_step_language(
        self,
        steps: list[SourceStep],
        expected: list[SourceStep],
        sheet: str,
        row: int,
        entity_id: str,
        headers: dict[str, int],
        issues: list[ValidationIssue],
    ) -> None:
        seen_focus_action = False
        for step in steps:
            if re.search(r"\b\d+\s*[,，]\s*\d+\b|坐标", step.text):
                issues.append(
                    self._issue(
                        "COORDINATE_FORBIDDEN",
                        Severity.ERROR,
                        "一期不允许在用例中使用绝对坐标",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(headers["steps"]),
                        entity_id=entity_id or None,
                        field="steps",
                    )
                )
            if "点击" in step.text:
                seen_focus_action = True
            if step.text.startswith("输入【") and not seen_focus_action:
                issues.append(
                    self._issue(
                        "INPUT_WITHOUT_FOCUS",
                        Severity.ERROR,
                        f"输入步骤 {step.ordinal} 前缺少可识别的输入框点击步骤",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(headers["steps"]),
                        entity_id=entity_id or None,
                        field="steps",
                    )
                )
        for result in expected:
            if result.text in {"页面正常", "操作成功", "结果正确"}:
                issues.append(
                    self._issue(
                        "EXPECTATION_NOT_OBSERVABLE",
                        Severity.ERROR,
                        f"预期结果“{result.text}”不可观察",
                        sheet=sheet,
                        row=row,
                        column=get_column_letter(headers["expected_results"]),
                        entity_id=entity_id or None,
                        field="expected_results",
                        suggestion="明确页面应显示或不显示的文案",
                    )
                )

    def _split_list(
        self,
        value: Any,
        separator: str,
        sheet: str,
        row: int,
        column: int,
        entity_id: str,
        field: str,
        issues: list[ValidationIssue],
    ) -> list[str]:
        text = self._text(value)
        if not text:
            return []
        if any(delimiter in text for delimiter in ("；", "，", ",")) and separator not in text:
            issues.append(
                self._issue(
                    "LIST_SEPARATOR_INVALID",
                    Severity.ERROR,
                    f"{field} 必须使用分隔符“{separator}”",
                    sheet=sheet,
                    row=row,
                    column=get_column_letter(column),
                    entity_id=entity_id or None,
                    field=field,
                )
            )
        values = [item.strip() for item in text.split(separator) if item.strip()]
        if len(values) != len(set(values)):
            issues.append(
                self._issue(
                    "LIST_VALUE_DUPLICATE",
                    Severity.ERROR,
                    f"{field} 包含重复值",
                    sheet=sheet,
                    row=row,
                    column=get_column_letter(column),
                    entity_id=entity_id or None,
                    field=field,
                )
            )
        return values

    @staticmethod
    def _parse_enabled(value: Any) -> bool:
        text = ExcelWorkbookImporter._text(value)
        if text == "是":
            return True
        if text == "否":
            return False
        raise ValueError("是否启用只接受“是”或“否”")

    @staticmethod
    def _parse_phase(value: Any) -> FixturePhase:
        mapping = {"前置": FixturePhase.SETUP, "后置": FixturePhase.TEARDOWN}
        text = ExcelWorkbookImporter._text(value)
        if text not in mapping:
            raise ValueError("阶段只接受“前置”或“后置”")
        return mapping[text]

    @staticmethod
    def _parse_scope(value: Any) -> FixtureScope:
        mapping = {
            "用例": FixtureScope.FUNCTION,
            "模块": FixtureScope.MODULE,
            "会话": FixtureScope.SESSION,
        }
        text = ExcelWorkbookImporter._text(value)
        if text not in mapping:
            raise ValueError("作用域只接受“用例”“模块”或“会话”")
        return mapping[text]

    @staticmethod
    def _parse_config_bool(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        text = str(value).strip().lower()
        if text == "true":
            return True
        if text == "false":
            return False
        raise ValueError("布尔值只接受 true 或 false")

    def _pydantic_issues(
        self,
        error: ValidationError,
        code: str,
        sheet: str,
        row: int,
        entity_id: str,
        headers: dict[str, int],
    ) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        for item in error.errors():
            field = str(item["loc"][0])
            column = get_column_letter(headers[field]) if field in headers else None
            issues.append(
                self._issue(
                    code,
                    Severity.ERROR,
                    f"{field} 不合法：{item['msg']}",
                    sheet=sheet,
                    row=row,
                    column=column,
                    entity_id=entity_id or None,
                    field=field,
                )
            )
        return issues

    @staticmethod
    def _snapshot(
        workbook_id: str,
        source_filename: str,
        source_sha256: str,
        cases: list[TestCaseDefinition],
        fixtures: list[FixtureDefinition],
        config: ProjectConfig | None,
        raw_config: dict[str, str | int | bool | None],
        issues: list[ValidationIssue],
        fixture_plans: dict[str, list[str]],
    ) -> WorkbookSnapshot:
        status = (
            WorkbookStatus.INVALID
            if any(issue.severity == Severity.ERROR for issue in issues)
            else WorkbookStatus.VALID
        )
        return WorkbookSnapshot(
            workbook_id=workbook_id,
            source_filename=source_filename,
            source_sha256=source_sha256,
            imported_at=datetime.now(UTC),
            status=status,
            cases=cases,
            fixtures=fixtures,
            config=config,
            raw_config=raw_config,
            validation_issues=issues,
            fixture_plans=fixture_plans,
        )

    @staticmethod
    def _case_ref_issue(
        code: str,
        case: TestCaseDefinition,
        field: str,
        column: str,
        message: str,
    ) -> ValidationIssue:
        return ValidationIssue(
            code=code,
            severity=Severity.ERROR,
            message=message,
            sheet="Cases",
            row=case.source_row,
            column=column,
            entity_id=case.case_id,
            field=field,
            suggestion="修正夹具 ID、启用状态或引用位置",
        )

    @staticmethod
    def _issue(
        code: str,
        severity: Severity,
        message: str,
        *,
        sheet: str | None = None,
        row: int | None = None,
        column: str | None = None,
        entity_id: str | None = None,
        field: str | None = None,
        suggestion: str | None = None,
    ) -> ValidationIssue:
        return ValidationIssue(
            code=code,
            severity=severity,
            message=message,
            sheet=sheet,
            row=row,
            column=column,
            entity_id=entity_id,
            field=field,
            suggestion=suggestion,
        )

    @staticmethod
    def _text(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _is_blank(value: Any) -> bool:
        return value is None or (isinstance(value, str) and not value.strip())
