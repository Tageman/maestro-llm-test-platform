"""Opt-in real Android MCP regression using isolated, safe navigation inputs.

Requires an online emulator and a visible unique navigation label. Never clears
App data, installs an APK, changes settings, or approves business user inputs.
Run with backend/src:backend/tests on PYTHONPATH; runtime evidence is private.
"""
from __future__ import annotations

import argparse
from datetime import UTC, datetime
import hashlib
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import yaml

from maestro_app.application.flow_translator import RuleBasedFlowTranslator
from maestro_app.application.generation_service import GenerationService
from maestro_app.application.workbook_service import WorkbookService
from maestro_app.cli.m3_verify import build_service
from maestro_app.core.config import Settings
from maestro_app.domain.models import ApprovalRequest, CreateGenerationRequest, CreateRunRequest
from maestro_app.domain.enums import RunMode, WorkbookStatus
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from maestro_app.infrastructure.mcp_diagnostics import MaestroMcpDiagnosticService
from maestro_app.main import create_app
from m1_helpers import build_sample_workbook


ROOT = Path(__file__).resolve().parents[1]
SCENARIOS = ("review-normal", "unattended-repair", "assertion-failure")


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def manifest():
    files = [ROOT / "pyproject.toml", ROOT / "frontend/package-lock.json"]
    for folder, pattern in (("backend/src", "*.py"), ("schemas", "*.json"), ("scripts", "*.py")):
        files.extend((ROOT / folder).rglob(pattern))
    return {str(path.relative_to(ROOT)): digest(path) for path in sorted(files)}


def scenario(root, name, args):
    root.mkdir()
    settings = Settings.from_environment().model_copy(update={
        "runtime_dir": root / "runtime", "android_device_id": args.device,
        "avd_name": args.avd, "app_id": args.app_id,
    })
    service = build_service(settings)
    source = build_sample_workbook(root / "source.xlsx")
    workbook = load_workbook(source)
    changes = {
        "app_id": args.app_id, "avd_name": args.avd, "android_device_id": args.device,
        "maestro_command": settings.maestro_command, "apk_path": None,
        "launch_clear_state": "false", "test_timeout_ms": 180000,
        "max_syntax_repair": 0, "max_locator_repair": 1, "max_total_run_attempts": 2,
        "output_dir": str(root / "export"),
    }
    for row in workbook["Config"].iter_rows(min_row=6):
        if row[0].value in changes:
            row[1].value = changes[row[0].value]
    fixtures = workbook["Fixtures"]
    fixtures.delete_rows(6, fixtures.max_row - 5)
    for identity, phase in (("S", "前置"), ("PS", "后置")):
        fixtures.append([identity, "只读导航检查", phase, "会话", "", "", 10, "是", f"1. 点击【{args.safe_text}】", f"1. 页面显示【{args.safe_text}】"])
    selector = args.safe_text + "旧" if name == "unattended-repair" else args.safe_text
    expected = "M4验收不存在的断言目标" if name == "assertion-failure" else args.safe_text
    cases = workbook["Cases"]
    cases.delete_rows(6, cases.max_row - 5)
    cases.append(["C1", "M4", "安全导航验收", "S", "PS", "P0", "acceptance", "是", f"1. 点击【{selector}】", f"1. 页面显示【{expected}】"])
    workbook.save(source)
    workbook.close()
    source_hash = digest(source)
    imported = WorkbookService(service.workbook_repository, ExcelWorkbookImporter()).import_bytes(source.name, source.read_bytes())
    assert imported.status == WorkbookStatus.VALID
    generation_service = GenerationService(service.workbook_repository, service.generation_repository, RuleBasedFlowTranslator())
    mode = RunMode.REVIEW if name == "review-normal" else RunMode.UNATTENDED
    generation = generation_service.create(CreateGenerationRequest(workbook_id=imported.workbook_id, mode=mode))
    if mode == RunMode.REVIEW:
        assert generation.status.value == "awaiting_review"
        # Harness approval of its own synthetic input, NOT independent human review.
        for flow in generation.flows:
            detail = generation_service.get_flow(generation.generation_id, flow.flow_id)
            assert detail.syntax_valid and not detail.warnings
        generation = generation_service.approve(generation.generation_id, ApprovalRequest(flow_ids=[flow.flow_id for flow in generation.flows]))
    assert generation.status.value == "ready"
    # Strict preflight proves real MCP format compatibility even without a repair.
    diagnostic = MaestroMcpDiagnosticService().diagnose(
        maestro_command=settings.maestro_command, device_id=args.device,
        failed_selector=args.safe_text + "旧", evidence_dir=root / "preflight",
    )
    assert diagnostic.candidate == args.safe_text, "Safe navigation label is not uniquely visible; prepare the app manually"
    save(root / "input.json", {"source_sha256": source_hash, "generation_id": generation.generation_id,
                               "mode": mode.value, "scenario": name, "device": args.device, "avd": args.avd})
    summary = service.create(CreateRunRequest(generation_id=generation.generation_id))
    run = service.run_repository.get_run(summary.run_id)
    run_root = service.run_repository.run_dir(run.run_id)
    save(root / "run-snapshot.json", run.model_dump(mode="json"))
    assert digest(source) == source_hash
    case = run.case_runs[0]
    assert case.cleanup_result == "passed"
    if name == "assertion-failure":
        assert run.state.value == "COMPLETED_WITH_FAILURES" and case.state.value == "FAILED"
        assert not run.repairs and len(case.attempts) == 1
        assert case.attempts[0].failure_type.value == "ASSERTION"
    else:
        assert run.state.value == "COMPLETED" and case.state.value == "PASSED"
        assert len(case.attempts) == (2 if name == "unattended-repair" else 1)
        assert len(run.repairs) == (1 if name == "unattended-repair" else 0)
    if name == "unattended-repair":
        assert run.repairs[0].after_selector == args.safe_text
        before = next(run_root.glob("repairs/**/before.yaml"))
        after = next(run_root.glob("repairs/**/after.yaml"))
        old = list(yaml.safe_load_all(before.read_text()))
        new = list(yaml.safe_load_all(after.read_text()))
        assert old[0] == new[0] and old[1][1:] == new[1][1:]
        assert len(old[1]) == len(new[1]) and "tapOn" in old[1][0] and "tapOn" in new[1][0]
    assert run.result_workbook_path
    result = run_root / run.result_workbook_path
    checked = load_workbook(result, read_only=True)
    checked.close()
    # Fresh application lifespan and API routes; no substitute for human UI review.
    with TestClient(create_app(settings)) as client:
        assert client.get(f"/api/v1/runs/{run.run_id}").json()["state"] == run.state.value
        response = client.get(f"/api/v1/runs/{run.run_id}/result-workbook")
        assert response.status_code == 200 and hashlib.sha256(response.content).hexdigest() == digest(result)
        for artifact in run.artifacts:
            response = client.get(f"/api/v1/runs/{run.run_id}/artifacts/{artifact.artifact_id}")
            assert response.status_code == 200
            assert hashlib.sha256(response.content).hexdigest() == artifact.sha256
    return {"status": "PASS", "run_id": run.run_id, "state": run.state.value,
            "attempts": len(case.attempts), "repairs": len(run.repairs),
            "source_sha256": source_hash, "result_sha256": digest(result),
            "artifact_downloads_verified": len(run.artifacts)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--device", required=True)
    parser.add_argument("--avd", required=True)
    parser.add_argument("--app-id", required=True)
    parser.add_argument("--safe-text", required=True)
    parser.add_argument("--scenario", action="append", choices=SCENARIOS)
    args = parser.parse_args()
    root = args.root.resolve()
    root.mkdir(parents=True, exist_ok=False)
    frozen = manifest()
    save(root / "code-manifest.json", frozen)
    results = []
    for name in args.scenario or SCENARIOS:
        print(f"START {name}", flush=True)
        try:
            record = {"scenario": name, **scenario(root / name, name, args)}
            assert manifest() == frozen, "Execution sources changed during acceptance"
        except Exception as exc:
            record = {"scenario": name, "status": "FAIL", "error": f"{type(exc).__name__}: {exc}"}
        results.append(record)
        save(root / "summary.json", {"tested_at": datetime.now(UTC).isoformat(), "results": results,
                                    "scope": "real-device synthetic delta regression, not business baseline or human signoff"})
        print(json.dumps(record, ensure_ascii=False), flush=True)
        if record["status"] != "PASS":
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
