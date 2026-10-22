"""Opt-in real-device crash acceptance. Never kills a pre-existing service/device.

Creates synthetic Excel input and a dedicated backend process per phase. Pauses
only after real adapter work, SIGKILLs that exact child, then verifies recovery.
Requires backend/src and backend/tests on PYTHONPATH and an already booted AVD.
Local evidence includes screenshots/hierarchy; never commit runtime/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import signal
import subprocess
import sys
import time
from datetime import UTC, datetime

from openpyxl import load_workbook

from maestro_app.application.flow_translator import RuleBasedFlowTranslator
from maestro_app.application.generation_service import GenerationService
from maestro_app.application.workbook_service import WorkbookService
from maestro_app.cli.m3_verify import build_service
from maestro_app.core.config import Settings
from maestro_app.domain.models import CreateGenerationRequest, CreateRunRequest
from maestro_app.domain.enums import RunMode, RunState, WorkbookStatus
from maestro_app.infrastructure.excel_importer import ExcelWorkbookImporter
from maestro_app.infrastructure.mcp_client import MaestroMcpClient
from maestro_app.infrastructure.execution_guard import RuntimeExecutionGuard
from maestro_app.application.errors import StateConflict
from maestro_app.application.run_recovery import RunRecovery
from maestro_app.infrastructure.run_repository import RunRepository
from m1_helpers import build_sample_workbook

ROOT = Path(__file__).resolve().parents[1]
PHASES = ("PREPARING_DEVICE", "RUNNING", "DIAGNOSING", "REPAIRING", "FINALIZING")


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def settings(root):
    return Settings.from_environment().model_copy(update={"runtime_dir": root / "runtime", "android_device_id": "emulator-5554"})


def pause(root, phase, detail):
    save(root / "pause.json", {"phase": phase, "pid": os.getpid(), "time": datetime.now(UTC).isoformat(), "detail": detail})
    while True:
        time.sleep(0.2)


def child(root, phase, text):
    cfg = settings(root)
    service = build_service(cfg)
    source = build_sample_workbook(root / "source.xlsx")
    workbook = load_workbook(source)
    changes = {
        "maestro_command": cfg.maestro_command, "android_device_id": cfg.android_device_id,
        "launch_clear_state": "false", "test_timeout_ms": 180000,
        "max_syntax_repair": 0, "max_locator_repair": 1, "max_total_run_attempts": 2,
        "output_dir": str(root / "export"),
    }
    for row in workbook["Config"].iter_rows(min_row=6):
        if row[0].value in changes:
            row[1].value = changes[row[0].value]
    fixtures = workbook["Fixtures"]
    fixtures.delete_rows(6, fixtures.max_row - 5)
    fixtures.append(["S", "会话只读检查", "前置", "会话", "", "", 10, "是", f"1. 点击【{text}】", f"1. 页面显示【{text}】"])
    fixtures.append(["PS", "会话后置检查", "后置", "会话", "", "", 20, "是", f"1. 点击【{text}】", f"1. 页面显示【{text}】"])
    cases = workbook["Cases"]
    cases.delete_rows(6, cases.max_row - 5)
    selector = text + "旧" if phase in {"RUNNING", "DIAGNOSING", "REPAIRING"} else text
    for case_id, module in [("A1", "A"), ("A2", "A"), ("B1", "B")]:
        cases.append([case_id, module, "只读导航验收", "S", "PS", "P0", "acceptance", "是", f"1. 点击【{selector}】", f"1. 页面显示【{text}】"])
    workbook.save(source)
    workbook.close()
    imported = WorkbookService(service.workbook_repository, ExcelWorkbookImporter()).import_bytes(source.name, source.read_bytes())
    assert imported.status == WorkbookStatus.VALID, imported
    generation = GenerationService(service.workbook_repository, service.generation_repository, RuleBasedFlowTranslator()).create(CreateGenerationRequest(workbook_id=imported.workbook_id, mode=RunMode.UNATTENDED))
    save(root / "input.json", {"source_sha256": digest(source), "generation_id": generation.generation_id, "phase": phase})

    if phase == "PREPARING_DEVICE":
        original = service.device_manager.prepare
        def prepare(config):
            result = original(config)
            pause(root, phase, result.model_dump(mode="json"))
        service.device_manager.prepare = prepare
    elif phase == "RUNNING":
        original = service.pytest_runner.run
        def run(run_id, package, **kwargs):
            tick = kwargs.get("on_tick")
            def on_tick():
                if tick:
                    tick()
                # Wait for the actual Maestro case command, not just Python's
                # pre-spawn trace. Excludes fixture/default-launch logs.
                folder = package.junit_path.parents[1] / "artifacts/maestro/A1/case"
                for log in folder.rglob("maestro.log"):
                    if "CliConsoleListener.onCommandStart" in log.read_text(errors="replace"):
                        pause(root, phase, {"command_log": str(log), "sha256": digest(log)})
            kwargs["on_tick"] = on_tick
            return original(run_id, package, **kwargs)
        service.pytest_runner.run = run
    elif phase == "DIAGNOSING":
        original = MaestroMcpClient.call_tool
        def call_tool(client, name, arguments=None):
            response = original(client, name, arguments)
            if name == "take_screenshot":
                save(root / "mcp-screenshot-response.json", response)
                pause(root, phase, {"tool": name, "response_sha256": digest(root / "mcp-screenshot-response.json"), "mcp_still_open": True})
            return response
        MaestroMcpClient.call_tool = call_tool
    elif phase == "REPAIRING":
        original = service.syntax_checker.check
        def check(path, **kwargs):
            result = original(path, **kwargs)
            if "repairs" in path.parts:
                assert result.valid, result.output
                pause(root, phase, {"repaired_yaml": str(path), "sha256": digest(path), "syntax_exit": result.exit_code})
            return result
        service.syntax_checker.check = check
    elif phase == "FINALIZING":
        original = service.result_writer.write
        def write(**kwargs):
            result = original(**kwargs)
            pause(root, phase, {"workbook": str(result.path), "sha256": digest(result.path)})
        service.result_writer.write = write
    result = service.create(CreateRunRequest(generation_id=generation.generation_id))
    save(root / "completed.json", result.model_dump(mode="json"))


def descendants(pid):
    rows = subprocess.check_output(["ps", "-axo", "pid=,ppid="], text=True)
    pairs = [tuple(map(int, line.split())) for line in rows.splitlines()]
    found = {pid}
    while True:
        added = {child for child, parent in pairs if parent in found} - found
        if not added:
            return sorted(found - {pid})
        found.update(added)


def alive(pid):
    try:
        os.kill(pid, 0)
        return True
    except ProcessLookupError:
        return False


def parent(root, phases, text):
    root.mkdir(parents=True, exist_ok=False)
    code = {str(path.relative_to(ROOT)): digest(path) for folder in ["backend/src", "scripts", "schemas"] for path in (ROOT / folder).rglob("*") if path.is_file() and "__pycache__" not in path.parts}
    save(root / "code-manifest.json", code)
    results = []
    for phase in phases:
        target = root / phase.lower()
        target.mkdir()
        print(f"START {phase}", flush=True)
        env = dict(os.environ, PYTHONPATH=os.pathsep.join([str(ROOT / "backend/src"), str(ROOT / "backend/tests")]))
        with (target / "backend.log").open("wb") as log:
            backend = subprocess.Popen([sys.executable, __file__, "--child", "--root", str(target), "--phase", phase, "--safe-text", text], cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
            owned = []
            record = {"phase": phase, "status": "FAIL"}
            try:
                deadline = time.monotonic() + 900
                while not (target / "pause.json").exists():
                    if backend.poll() is not None:
                        raise RuntimeError(f"backend ended before phase marker: {backend.returncode}")
                    if time.monotonic() > deadline:
                        raise TimeoutError("phase not reached within 900s")
                    time.sleep(0.25)
                repository = RunRepository(settings(target).runtime_dir)
                run = repository.list_runs()[0]
                assert run.state.value == phase, run.state
                owned = descendants(backend.pid)
                record.update({"run_id": run.run_id, "backend_pid": backend.pid, "owned_pids": owned, "before_state": run.state.value})
                backend.send_signal(signal.SIGKILL)
                backend.wait(timeout=10)
                guard = RuntimeExecutionGuard(settings(target).runtime_dir)
                lease = None
                deadline = time.monotonic() + 30
                while lease is None and time.monotonic() < deadline:
                    try:
                        lease = guard.acquire()
                    except StateConflict:
                        time.sleep(0.1)
                assert lease is not None, "guardian did not release runtime lease"
                try:
                    deadline = time.monotonic() + 10
                    while any(alive(pid) for pid in owned) and time.monotonic() < deadline:
                        time.sleep(0.1)
                    assert not any(alive(pid) for pid in owned), "owned child survived backend exit"
                    run_root = repository.run_dir(run.run_id)
                    # Recovery appends audit artifacts to the ledger. Original
                    # entries and every execution payload must remain intact.
                    ledger = run_root / "result/artifacts.json"
                    before_ledger = json.loads(ledger.read_text()) if ledger.exists() else []
                    immutable = {str(path.relative_to(run_root)): digest(path) for path in run_root.rglob("*") if path.is_file() and str(path.relative_to(run_root)) not in {"run.json", "result/artifacts.json"}}
                    recovery = RunRecovery(repository)
                    assert recovery.recover() == [run.run_id]
                    assert recovery.recover() == []
                    after = repository.get_run(run.run_id)
                    assert after.state == RunState.FAILED, after.state
                    assert all(digest(run_root / path) == sha for path, sha in immutable.items()), "recovery rewrote execution evidence"
                    after_ledger = json.loads(ledger.read_text())
                    assert all(item in after_ledger for item in before_ledger), "recovery changed an existing artifact entry"
                    assert all(item["kind"] in {"recovery_report", "recovery_snapshot"} for item in after_ledger if item not in before_ledger), "unexpected new execution artifact"
                    assert digest(target / "source.xlsx") == json.loads((target / "input.json").read_text())["source_sha256"]
                finally:
                    lease.close()
                # Exercise normal service/API read and evidence download after
                # restart, without starting another run.
                from fastapi.testclient import TestClient
                from maestro_app.main import create_app
                with TestClient(create_app(settings(target), run_service=build_service(settings(target)))) as client:
                    summary = client.get(f"/api/v1/runs/{run.run_id}").json()
                    artifact_id = summary["recovery_report_artifact_id"]
                    response = client.get(f"/api/v1/runs/{run.run_id}/artifacts/{artifact_id}")
                    assert response.status_code == 200
                    assert response.json()["policy"] == "safe-stop-no-replay"
                record.update({"status": "PASS", "after_state": after.state.value, "no_surviving_owned_children": True, "no_evidence_replay_or_rewrite": True, "api_recovery_download": True})
            except Exception as exc:
                record["error"] = f"{type(exc).__name__}: {exc}"
            finally:
                if backend.poll() is None:
                    backend.kill()  # Only the exact child created above.
                backend.wait(timeout=10)
                save(target / "evidence.json", record)
                results.append(record)
                save(root / "summary.json", {"tested_at": datetime.now(UTC).isoformat(), "scope": "real-device safe synthetic scenarios, not business AC/human signoff", "results": results})
        print(f"{record['status']} {phase}: {record.get('error', '')}", flush=True)
        if record["status"] != "PASS":
            return 1
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", required=True, type=Path)
    parser.add_argument("--safe-text", required=True, help="Visible, unique, safe navigation label; never submit/delete/pay")
    parser.add_argument("--phase", choices=PHASES, action="append")
    parser.add_argument("--child", action="store_true")
    args = parser.parse_args()
    if args.child:
        child(args.root.resolve(), args.phase[0], args.safe_text)
    else:
        raise SystemExit(parent(args.root.resolve(), args.phase or PHASES, args.safe_text))
